"""Assessment pack: quiz bank workbook with answer key, rubric and evaluation sheets."""
import pandas as pd
from pathlib import Path
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = Path("assessment"); OUT.mkdir(exist_ok=True)
NAVY="0B2545"; MID="13497B"; TEAL="1C7293"; PAPER="F1F5F9"; GOLD="D9A441"
THIN=Side(style="thin", color="D8E0E8")
BORD=Border(left=THIN,right=THIN,top=THIN,bottom=THIN)

QUIZ = [
 (1,"M1","Why can't you shuffle a time series before splitting?",
  "Order is the signal; shuffling leaks the future into the past — test points end up preceding training points.","Recall"),
 (2,"M1",'What does `asfreq("h")` do that `parse_dates` alone does not?',
  "It declares the frequency and materialises missing timestamps as NaN, turning a silent gap into a visible one.","Recall"),
 (3,"M1","Additive or multiplicative — which for a seasonal swing that grows with the level?",
  "Multiplicative (or take logs and model additively).","Apply"),
 (4,"M2","ADF p = 0.01 and KPSS p = 0.01 — what is the reconciled verdict?",
  "Difference-stationary. ADF says stationary, KPSS disagrees, so difference once and re-test.","Analyse"),
 (5,"M2","For seasonal data, seasonal or regular differencing first?",
  "Seasonal first, then re-test; apply a regular difference only if it is still non-stationary.","Apply"),
 (6,"M2","ACF tails off, PACF cuts off after lag 2 — model and order?","AR(2).","Analyse"),
 (7,"M3","What is m in SARIMA(p,d,q)(P,D,Q)m for a daily cycle on hourly data?",
  "24 — m is observations per cycle.","Recall"),
 (8,"M3","Why must SARIMAX exogenous regressors be known at forecast time?",
  "You need their future values to produce the forecast; using realised values instead is a leak that looks like skill.","Understand"),
 (9,"M3","A Ljung-Box p-value below 0.05 on the residuals means what?",
  "Residual autocorrelation remains — the model is under-specified; raise an order and refit.","Analyse"),
 (10,"M4","Why can a gradient-boosted tree never forecast a record-high demand?",
  "Trees interpolate within the training range; they cannot extrapolate beyond any value they have seen.","Understand"),
 (11,"M4","What does `.shift(1)` before a rolling feature prevent?",
  "The window including the target hour — a self-leak.","Apply"),
 (12,"M5","Prediction interval or confidence interval for a future observation?",
  "Prediction interval — wider, because it includes the irreducible noise as well as parameter uncertainty.","Recall"),
 (13,"M5","Minimising the pinball loss at tau = 0.9 recovers what?","The 90th percentile.","Understand"),
 (14,"M5","What does split-conformal guarantee, and under what assumption?",
  "Marginal (1 - alpha) coverage, under exchangeability of the calibration and test nonconformity scores.","Understand"),
 (15,"M5","Two intervals both cover 90% — which one ships?","The sharper (narrower) one.","Apply"),
 (16,"M6","Expanding versus sliding window — which one forgets old data?",
  "Sliding (the fixed-length window drops the oldest data as it advances).","Recall"),
 (17,"M6","What does MASE < 1 mean?","The model beats the seasonal-naive baseline.","Recall"),
 (18,"M6","Why does MAPE fail near zero and MASE not?",
  "MAPE's denominator goes to zero and the percentage explodes; MASE divides by a fixed in-sample baseline MAE.","Analyse"),
 (19,"M6","A non-significant Diebold-Mariano test tells you to do what?",
  "Ship the simpler or cheaper model — the accuracy gap is noise you cannot distinguish.","Evaluate"),
 (20,"M7","Marginal versus conditional coverage — why can a model be safe on one and dangerous on the other?",
  "90% on average can hide roughly 85% on the high-demand afternoons that are the only hours reserve is actually sized for.","Evaluate"),
]

def style_header(ws, ncols, row=1, fill=NAVY):
    for c in range(1, ncols+1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(bold=True, color="FFFFFF", size=11)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28
    ws.freeze_panes = ws.cell(row=row+1, column=1)

def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x

with pd.ExcelWriter(OUT/"SDA-DSC-212_Assessment_Pack.xlsx", engine="openpyxl") as xl:
    # ---------------------------------------------------------------- quiz --
    q = pd.DataFrame(QUIZ, columns=["#","Module","Question","Model answer","Bloom level"])
    q["Marks"]=1
    q.to_excel(xl, sheet_name="Quiz bank (answer key)", index=False)
    ws = xl.book["Quiz bank (answer key)"]
    style_header(ws, 6); widths(ws, [5,10,62,72,14,8])
    for r in range(2, len(q)+2):
        for c in range(1,7):
            ws.cell(row=r,column=c).alignment=Alignment(wrap_text=True, vertical="top")
            ws.cell(row=r,column=c).border=BORD
        ws.cell(row=r,column=4).font=Font(color="1E7A4D")
        ws.row_dimensions[r].height=34

    # student sheet — questions only
    qs = q[["#","Module","Question","Marks"]].copy(); qs["Your answer"]=""
    qs.to_excel(xl, sheet_name="Quiz (student sheet)", index=False)
    ws = xl.book["Quiz (student sheet)"]; style_header(ws,5,fill=MID); widths(ws,[5,10,70,8,60])
    for r in range(2,len(qs)+2):
        for c in range(1,6):
            ws.cell(row=r,column=c).alignment=Alignment(wrap_text=True,vertical="top")
            ws.cell(row=r,column=c).border=BORD
        ws.row_dimensions[r].height=32

    # --------------------------------------------------------- capstone rubric
    rub = pd.DataFrame([
     ["Data, index & decomposition",12,"Correct tz-aware frequency, explicit gap policy, MSTL with justified strengths and a justified additive/multiplicative choice","Minor gaps: implicit frequency, or an unjustified transform choice","Broken index, silent gaps, or the wrong `period`"],
     ["Stationarity & modelling rigour",15,"ADF + KPSS reconciled, minimal differencing, at least three families all leak-safe and diagnosed","One family thin, or residuals left undiagnosed","Leakage present, no baseline, or AIC compared across different d"],
     ["Feature engineering (leakage-safe)",15,"assert_no_leakage PASS; lags chosen from the ACF; .shift(1) rolling; Fri-Sat weekend and Hijri calendar correct","One subtle leak caught late, or the weekend miscoded","A shuffle split or a centred window survives; future exogenous values used"],
     ["Probabilistic calibration",15,"Coverage 0.88-0.92 marginal AND conditional; CQR justified; sharpest interval at the target coverage","Marginal coverage in band but the conditional gap is unaddressed","A confidence interval where a prediction interval is required; coverage off target"],
     ["Backtesting validity",18,"At least 50 origins, training data ending at or before each origin, MASE + coverage + pinball + per-origin breakdown","Fewer origins, or the probabilistic metrics are missing","Fit-once or in-sample scoring; a single split; no baseline"],
     ["Model selection & evidence",15,"A DM-significant choice with constraints and weights stated, plus flip-conditions, a fallback and a monitoring contract","The choice is reasonable but the evidence or the stated limits are thin","Lowest-MAE pick with no significance test and no stated limits"],
     ["Report, brief & reproducibility",10,"One-command reproduction; a decision-language brief; an auditor-ready report","Runs with some fiddling; the brief is partly jargon","Irreproducible; jargon-only; the numbers cannot be regenerated"],
    ], columns=["Criterion","Weight","90-100% band","70-89% band","Below 70% band"])
    rub.loc[len(rub)] = ["TOTAL", rub.Weight.sum(), "Pass >= 70 · Distinction >= 90",
                         "One chosen extension adds up to +5 (capped at 100), and only if mandatory scope scores >= 80","" ]
    rub.to_excel(xl, sheet_name="Capstone rubric",index=False)
    ws=xl.book["Capstone rubric"]; style_header(ws,5); widths(ws,[34,9,58,50,54])
    for r in range(2,len(rub)+2):
        for c in range(1,6):
            ws.cell(row=r,column=c).alignment=Alignment(wrap_text=True,vertical="top")
            ws.cell(row=r,column=c).border=BORD
        ws.row_dimensions[r].height=52
    last=len(rub)+1
    for c in range(1,6):
        ws.cell(row=last,column=c).font=Font(bold=True,color=NAVY)
        ws.cell(row=last,column=c).fill=PatternFill("solid",fgColor=PAPER)

    # ------------------------------------------------------- scoring sheet ---
    sc = pd.DataFrame({"Participant":[f"(name {i})" for i in range(1,13)]})
    for crit, w in zip(rub.Criterion[:-1], rub.Weight[:-1]):
        sc[f"{crit} /{w}"] = ""
    sc["Mandatory subtotal"]=""; sc["Extension bonus (0-5)"]=""; sc["TOTAL /100"]=""
    sc["Outcome"]=""
    sc.to_excel(xl, sheet_name="Capstone scoring sheet",index=False)
    ws=xl.book["Capstone scoring sheet"]; style_header(ws,len(sc.columns),fill=TEAL)
    widths(ws,[22]+[15]*7+[18,18,14,14])
    n=len(sc)+1
    for r in range(2,n+1):
        # subtotal = sum of the seven criterion columns
        ws.cell(row=r,column=9).value=f"=SUM(B{r}:H{r})"
        ws.cell(row=r,column=11).value=f"=MIN(100, I{r} + IF(I{r}>=80, J{r}, 0))"
        ws.cell(row=r,column=12).value=(f'=IF(K{r}="","",IF(K{r}>=90,"Distinction",'
                                        f'IF(K{r}>=70,"Pass","Not yet")))')
        for c in range(1,13):
            ws.cell(row=r,column=c).border=BORD
        for c in (9,11,12): ws.cell(row=r,column=c).font=Font(bold=True,color=NAVY)
    ws.cell(row=n+2,column=1).value=("Bonus applies only when the mandatory subtotal is at least 80. "
        "Badge issuance also requires the overall course mark >= 70 and zero academic-integrity flags.")
    ws.cell(row=n+2,column=1).font=Font(italic=True,color="5A6B7C")

    # --------------------------------------------------- course weighting ----
    cw = pd.DataFrame([
     ["Lab completion (7 labs)",30,"Checkpoint commits plus the expected outputs for each lab"],
     ["PA-1 (Day 1) + PA-2 (Day 2)",20,"Artefacts plus written diagnosis notes"],
     ["Quiz (10 of the 20 questions)",10,"15 minutes, closed book"],
     ["Capstone",40,"The rubric above, graded repository-first"],
    ], columns=["Component","Weight %","Evidence"])
    cw.loc[len(cw)]=["TOTAL",100,"Pass >= 70 overall AND capstone >= 70 AND zero integrity flags"]
    cw.to_excel(xl, sheet_name="Course weighting",index=False)
    ws=xl.book["Course weighting"]; style_header(ws,3); widths(ws,[34,12,70])
    for r in range(2,len(cw)+2):
        for c in range(1,4):
            ws.cell(row=r,column=c).alignment=Alignment(wrap_text=True,vertical="center")
            ws.cell(row=r,column=c).border=BORD

    # ------------------------------------------------------- PA rubrics -----
    pa = pd.DataFrame([
     ["PA-1","Day 1","30 min","Index correctness",40,"Monotonic, unique, tz-aware, explicit frequency; row count before and after reported"],
     ["PA-1","Day 1","30 min","Gap policy",30,"Every missing hour surfaced; short gaps imputed and flagged; long gaps left flagged, not silently filled"],
     ["PA-1","Day 1","30 min","Decomposition + strengths",30,"Correct period; MSTL panel; daily and weekly strengths reported and read aloud correctly"],
     ["PA-2","Day 2","30 min","Leak diagnosis notes",40,"Both planted leaks named precisely — the shuffle split and the centred rolling window"],
     ["PA-2","Day 2","30 min","Correct fixes",40,"Time-ordered split; .shift(1) before rolling with center=False; assert_no_leakage passes"],
     ["PA-2","Day 2","30 min","Honest re-measurement",20,"The post-fix accuracy is reported without excuse, and the gap to the leaked number is quantified"],
    ], columns=["Assessment","When","Duration","Criterion","Weight %","Evidence of full marks"])
    pa.to_excel(xl, sheet_name="PA-1 and PA-2 rubrics",index=False)
    ws=xl.book["PA-1 and PA-2 rubrics"]; style_header(ws,6,fill=MID); widths(ws,[12,9,10,30,10,72])
    for r in range(2,len(pa)+2):
        for c in range(1,7):
            ws.cell(row=r,column=c).alignment=Alignment(wrap_text=True,vertical="top")
            ws.cell(row=r,column=c).border=BORD
        ws.row_dimensions[r].height=34

    # ------------------------------------------------- benchmarks template --
    bm = pd.DataFrame([
     ["M1","Materialised gaps identified","100% of missing hours surfaced","count of NaN after asfreq",""],
     ["M1","Index correctness","monotonic, unique, tz-aware, freq set","df.index checks",""],
     ["M1","Seasonal strength (daily)","reported in [0,1]","Hyndman strength formula","0.97"],
     ["M1","Seasonal strength (weekly)","reported in [0,1]","Hyndman strength formula","0.61"],
     ["M2","Stationarity achieved","ADF p < 0.05 and KPSS p > 0.05","stationarity_report verdict","D=1, d=0"],
     ["M2","Differencing minimality","lowest (d,D) that passes","lag-1 ACF not strongly negative",""],
     ["M3","Residual white noise","Ljung-Box p > 0.05","residual_diagnostics","report honestly"],
     ["M3","Day-ahead MAPE (daily peak)","<= 2.5%","28-day hold-out","2.05%"],
     ["M3","Value of the temperature regressor","MAE delta vs no-exog","two fits compared","20 MW"],
     ["M4","Leakage check","PASS, 0 future-referencing features","assert_no_leakage","PASS (44)"],
     ["M4","Day-ahead hourly MASE","< 0.7","rolling origins","0.445"],
     ["M5","Empirical coverage (marginal)","0.88 - 0.92","held-out window","0.902"],
     ["M5","Empirical coverage (peak hours)","> = 0.88","conditional report","0.845 — FLAG"],
     ["M6","Number of rolling origins",">= 50","backtest config","60"],
     ["M6","Champion beats baseline","DM p < 0.05 vs the runner-up","diebold_mariano","p < 0.001"],
     ["M7","Champion changes across scenarios","at least 2 of 3 differ","constraint-weighted scoring","A/B/C differ"],
    ], columns=["Module","Metric","Target","How measured","Reference run"])
    bm["YOUR RESULT"]=""
    bm.to_excel(xl, sheet_name="BENCHMARKS template",index=False)
    ws=xl.book["BENCHMARKS template"]; style_header(ws,6,fill=TEAL); widths(ws,[9,36,32,34,18,20])
    for r in range(2,len(bm)+2):
        for c in range(1,7):
            ws.cell(row=r,column=c).alignment=Alignment(wrap_text=True,vertical="center")
            ws.cell(row=r,column=c).border=BORD
        ws.cell(row=r,column=6).fill=PatternFill("solid",fgColor="FFF6E0")

print("wrote", OUT/"SDA-DSC-212_Assessment_Pack.xlsx")
