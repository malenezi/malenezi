#!/usr/bin/env python3
"""
build_layer_b.py -- generate the synthetic internal overlay of the Dalil corpus.

SDA-AIE-214 -- Building RAG Systems, SDAIA Academy.

Produces, deterministically from a seed, the four things a real public
regulatory corpus cannot give a RAG course:

  * access-tiered documents          -> Module 5 access filter has teeth
  * superseded/current document pairs-> the corpus_v2 staleness drill (F8)
  * image-only Arabic pages          -> the OCR path and CER gate (F1)
  * tabular entitlement facts        -> the "table nobody could retrieve" (F1)

Outputs
-------
  data/corpus/corpus_v1/          documents as shipped on Day 1
  data/corpus/corpus_v2/          Day-4 upgrade: superseding editions land here
  data/corpus/ocr_gold/           clean transcripts of the scanned pages (CER)
  data/corpus/layer_b_index.json  ground truth: every doc, tier, lifecycle, facts

The `facts` recorded in layer_b_index.json are what make automatic golden-set
construction possible: build_golden_set.py reads them, so answer keys can never
drift from the corpus that was actually generated.

Usage
-----
    python corpus/build_layer_b.py
    python corpus/build_layer_b.py --out data/corpus --spec corpus/layer_b_specs/layer_b.yaml
"""
from __future__ import annotations

import argparse
import json
import random
import re
import sys
from datetime import date, timedelta
from pathlib import Path

import yaml

sys.path.insert(0, str(Path(__file__).resolve().parent))
from content_blocks import (  # noqa: E402
    ORG_EN, ORG_AR, LAYER_A_REFS, HR_TOPICS, IT_TOPICS, CIRCULAR_SUBJECTS,
    INTRANET_TOPICS, BOILERPLATE_HTML_NAV, BOILERPLATE_HTML_FOOTER,
)

ROOT = Path(__file__).resolve().parent.parent
FONT = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"

# Arabic rasterisation for the "scanned" family. Order matters: the first font
# found wins. A real Arabic typeface plus PIL's Raqm/HarfBuzz layout engine is
# what makes the scanned pages OCR-able at a realistic 2-5% CER; without it the
# generator falls back to arabic_reshaper + bidi, whose presentation-form
# glyphs Tesseract reads at ~20% CER and in VISUAL order -- a page no CER gate
# could ever pass. If your build prints the fallback warning, install
# fonts-noto-core (Debian/Ubuntu) or point AR_FONT_CANDIDATES at a Naskh face.
AR_FONT_CANDIDATES = [
    "/usr/share/fonts/truetype/noto/NotoNaskhArabic-Regular.ttf",
    "/usr/share/fonts/truetype/noto/NotoNaskhArabic-Bold.ttf",
    "/usr/share/fonts/opentype/fonts-hosny-amiri/Amiri-Regular.ttf",
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
]


def _ar_font_path() -> str:
    from pathlib import Path as _P
    for c in AR_FONT_CANDIDATES:
        if _P(c).exists():
            return c
    return FONT


def _raqm_available() -> bool:
    try:
        from PIL import features
        return bool(features.check("raqm"))
    except Exception:
        return False

DISC_EN = "FICTIONAL TRAINING DOCUMENT - SDAIA Academy SDA-AIE-214. Not a real instrument."
DISC_AR = "وثيقة تدريبية افتراضية - أكاديمية سدايا SDA-AIE-214. ليست وثيقة رسمية."


# ---------------------------------------------------------------- utilities
def ar(text: str) -> str:
    """Shape + bidi-order Arabic for raster rendering."""
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
        return get_display(arabic_reshaper.reshape(text))
    except ImportError:
        return text


AR_RUN = None


def fix_ar_runs(text: str) -> str:
    """Reshape + bidi-order any Arabic run inside an otherwise-Latin string.

    ReportLab has no bidi engine: handing it logical-order Arabic produces a
    PDF whose glyphs run left-to-right, i.e. visually reversed. Passing the
    visually-ordered presentation forms instead produces a correct-looking page
    whose EXTRACTED text is Arabic presentation forms (U+FB50-U+FEFF) in
    logical order.

    That is not a bug we are papering over -- it is the single most common
    Arabic-PDF pathology in real government document sets, and dalil.arabic
    .normalise_arabic() is the course's answer to it (NFKC folding). Ship it.
    """
    global AR_RUN
    if AR_RUN is None:
        import re as _re
        AR_RUN = _re.compile(r"[\u0600-\u06FF][\u0600-\u06FF\s\u060C\u061B\u061F.,:()\-/0-9\u0660-\u0669]*")
    return AR_RUN.sub(lambda m: ar(m.group(0)), text)


def hijri_ish(g: date) -> str:
    """Approximate Hijri year label -- close enough to teach digit/calendar
    normalisation without pretending to be a conversion library."""
    return f"{int((g.year - 622) * 33 / 32)}هـ"


ARABIC_DIGITS = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")


def to_arabic_digits(s: str) -> str:
    return s.translate(ARABIC_DIGITS)


# ------------------------------------------------------------------ writers
def write_digital_pdf(path: Path, title: str, subtitle: str, blocks: list, footer: str,
                      rtl: bool = False) -> None:
    """A clean, text-layer PDF. `blocks` is a list of ('h'|'p'|'table', payload)."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle)

    for name, f in (("DJV", FONT), ("DJV-B", FONT_BOLD)):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, f))

    ss = getSampleStyleSheet()
    align = 2 if rtl else 0
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName="DJV-B", fontSize=15,
                        leading=20, alignment=align)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontName="DJV-B", fontSize=11.5,
                        leading=16, alignment=align, spaceBefore=10)
    body = ParagraphStyle("body", parent=ss["BodyText"], fontName="DJV", fontSize=9.5,
                          leading=14, alignment=align)
    small = ParagraphStyle("small", parent=body, fontSize=7.5, textColor=colors.grey)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(path), pagesize=A4, title=title,
                            author=ORG_EN, subject="SDA-AIE-214 synthetic training corpus",
                            leftMargin=20 * mm, rightMargin=20 * mm,
                            topMargin=18 * mm, bottomMargin=18 * mm)
    title = fix_ar_runs(title)
    flow = [Paragraph(title, h1)]
    if subtitle:
        flow.append(Paragraph(fix_ar_runs(subtitle), small))
    flow.append(Spacer(1, 6))
    for kind, payload in blocks:
        if kind == "h":
            flow.append(Paragraph(fix_ar_runs(payload), h2))
        elif kind == "p":
            flow.append(Paragraph(fix_ar_runs(payload), body))
            flow.append(Spacer(1, 4))
        elif kind == "table":
            payload = [[fix_ar_runs(str(c)) for c in row] for row in payload]
            t = Table(payload, hAlign="LEFT", repeatRows=1)
            t.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), "DJV"),
                ("FONTNAME", (0, 0), (-1, 0), "DJV-B"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#8a8a8a")),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef2")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ]))
            flow.append(t)
            flow.append(Spacer(1, 8))
    flow.append(Spacer(1, 10))
    flow.append(Paragraph(footer, small))
    doc.build(flow)



AR_CHARS = re.compile(r"[\u0600-\u06FF]")


def _is_latin_line(text: str) -> bool:
    ar_n = len(AR_CHARS.findall(text))
    latin_n = sum(1 for c in text if c.isascii() and c.isalnum())
    return latin_n > ar_n


def _measure(draw, text, font, shaped: bool) -> float:
    try:
        if shaped:
            return draw.textlength(text, font=font, direction="rtl", language="ar")
    except Exception:
        pass
    return draw.textlength(text, font=font)


def _wrap_rtl(draw, text: str, font, max_w: float) -> list[str]:
    """Greedy word wrap measured with the same layout engine used to draw."""
    words = text.split()
    if not words:
        return []
    lines, cur = [], words[0]
    for w in words[1:]:
        cand = f"{cur} {w}"
        if _measure(draw, cand, font, True) > max_w:
            lines.append(cur)
            cur = w
        else:
            cur = cand
    lines.append(cur)
    return lines


def _draw_line(draw, text: str, font, x_right: float, y: float, rng, shaped: bool) -> None:
    fill = rng.randint(10, 55)
    if shaped:
        draw.text((x_right, y), text, font=font, fill=fill,
                  direction="rtl", language="ar", anchor="ra")
    else:
        txt = ar(text)
        draw.text((x_right - draw.textlength(txt, font=font), y), txt, font=font, fill=fill)


def write_scanned_pdf(path: Path, lines: list[str], seed: int) -> None:
    """Render Arabic text to noisy raster pages -> no text layer at all."""
    from PIL import Image, ImageDraw, ImageFont, ImageFilter

    rng = random.Random(seed)
    W, H = 1654, 2339                       # A4 @ 200 dpi
    ar_font_path = _ar_font_path()
    shaped = _raqm_available()
    if not shaped:
        print("  [warn] PIL Raqm layout unavailable -- scanned pages will OCR poorly "
              "(install libraqm / a Pillow wheel with Raqm support)")
    font = ImageFont.truetype(ar_font_path, 40)
    font_h = ImageFont.truetype(ar_font_path, 50)

    pages, cur, y = [], [], 0
    per_page = 22
    for i in range(0, len(lines), per_page):
        pages.append(lines[i:i + per_page])

    images = []
    for pg in pages:
        im = Image.new("L", (W, H), 252)
        d = ImageDraw.Draw(im)
        y = 120
        for idx, raw in enumerate(pg):
            f = font_h if raw.startswith("#") else font
            src = raw.lstrip("# ").strip()
            if not src:
                y += 22
                continue
            # Wrap to the page box. Without this, long Arabic lines anchored to
            # the right margin run off the LEFT edge of the page and the tail of
            # every long sentence is simply missing from the raster -- which
            # shows up later as an OCR "error" that is really a rendering bug.
            max_w = W - 150 - 110
            latin_line = _is_latin_line(src)
            for piece in _wrap_rtl(d, src, f, max_w):
                if latin_line:
                    # A predominantly-Latin line (document numbers, dates,
                    # course codes) laid out RTL confuses both the renderer and
                    # the OCR engine. Real bilingual documents set these LTR.
                    d.text((150, y), piece, font=f, fill=rng.randint(10, 55))
                else:
                    _draw_line(d, piece, f, W - 150 + rng.randint(-2, 2), y, rng, shaped)
                y += 64 if f is font_h else 56
            continue
        # scanner artefacts: speckle, slight rotation, blur, edge shadow
        px = im.load()
        for _ in range(int(W * H * 0.0006)):
            px[rng.randrange(W), rng.randrange(H)] = rng.randint(90, 190)
        im = im.rotate(rng.uniform(-0.5, 0.5), resample=Image.BILINEAR, fillcolor=252)
        im = im.filter(ImageFilter.GaussianBlur(0.35))
        d2 = ImageDraw.Draw(im)
        d2.rectangle([0, 0, 18, H], fill=228)
        images.append(im.convert("RGB"))

    path.parent.mkdir(parents=True, exist_ok=True)
    # Build the PDF from PNG rasters via PyMuPDF: no text layer is written, so
    # every extractor must fall through to OCR -- which is the point.
    import io
    import pymupdf

    out = pymupdf.open()
    for im in images:
        buf = io.BytesIO()
        im.save(buf, format="PNG")
        buf.seek(0)
        rect = pymupdf.Rect(0, 0, W * 72 / 200, H * 72 / 200)
        page = out.new_page(width=rect.width, height=rect.height)
        page.insert_image(rect, stream=buf.getvalue())
    out.set_metadata({"title": path.stem, "author": ORG_EN,
                      "subject": "X-Dalil-Synthetic scanned page (SDA-AIE-214)"})
    out.save(str(path), deflate=True)
    out.close()


def write_docx(path: Path, title: str, sections: list[tuple[str, list[str]]], footer: str) -> None:
    from docx import Document
    from docx.shared import Pt

    path.parent.mkdir(parents=True, exist_ok=True)
    d = Document()
    d.core_properties.title = title
    d.core_properties.author = ORG_EN
    d.core_properties.comments = "X-Dalil-Synthetic: SDA-AIE-214 training corpus"
    d.add_heading(title, level=1)
    for head, paras in sections:
        d.add_heading(head, level=2)
        for p in paras:
            if p.startswith("- "):
                d.add_paragraph(p[2:], style="List Bullet")
            else:
                d.add_paragraph(p)
    fp = d.add_paragraph(footer)
    fp.runs[0].font.size = Pt(7.5)
    d.save(path)


def write_xlsx(path: Path, sheet_title: str, header: list[str], rows: list[list],
               notes: list[str]) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    thin = Side(style="thin", color="8A8A8A")
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="E8EEF2")
        c.alignment = Alignment(horizontal="center", wrap_text=True)
        c.border = Border(thin, thin, thin, thin)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = Border(thin, thin, thin, thin)
    for i, w in enumerate([28, 20, 20, 22, 34][: len(header)], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    ws.append([])
    for n in notes:
        ws.append([n])
    wb.save(path)


def write_html(path: Path, title_en: str, title_ar: str, body_en: list[str],
               body_ar: list[str], crumb: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    en = "\n".join(f"<p>{p}</p>" for p in body_en)
    arb = "\n".join(f"<p>{p}</p>" for p in body_ar)
    html = f"""<!doctype html>
<html lang="en"><head><meta charset="utf-8">
<title>{title_en} — {ORG_EN} intranet</title>
<meta name="x-dalil-synthetic" content="SDA-AIE-214">
<style>body{{font-family:system-ui;max-width:820px;margin:2rem auto}}</style>
</head><body>
{BOILERPLATE_HTML_NAV.format(crumb=crumb)}
<main>
<h1>{title_en}</h1>
{en}
<hr>
<section dir="rtl" lang="ar"><h2>{title_ar}</h2>
{arb}
</section>
</main>
{BOILERPLATE_HTML_FOOTER}
</body></html>"""
    path.write_text(html, encoding="utf-8")


# --------------------------------------------------------------- generation
class Builder:
    def __init__(self, spec: dict, out: Path):
        self.spec = spec
        self.out = out
        self.rng = random.Random(spec.get("seed", 20260909))
        self.v1 = out / "corpus_v1"
        self.v2 = out / "corpus_v2"
        self.gold = out / "ocr_gold"
        self.index: list[dict] = []
        self.base_date = date(2024, 1, 15)

    # -- helpers -----------------------------------------------------------
    def record(self, **kw):
        self.index.append(kw)

    def a_ref(self):
        return self.rng.choice(LAYER_A_REFS)

    def eff_date(self, offset_days: int) -> date:
        return self.base_date + timedelta(days=offset_days)

    # -- families ----------------------------------------------------------
    def hr_scanned_ar(self, count: int):
        made = 0
        for topic_id, t_en, t_ar, desc_en, desc_ar in HR_TOPICS:
            for variant in range(2):
                if made >= count:
                    return
                made += 1
                n = made
                doc_id = f"ndsa-hr-{topic_id}-ar-{variant+1}"
                ref_id, ref_en, ref_ar = self.a_ref()
                eff = self.eff_date(30 * n)
                pct = 15 + (n * 3) % 25
                days = 20 + (n * 2) % 15
                lines = [
                    f"# {ORG_AR}",
                    f"# {t_ar}",
                    f"رقم الوثيقة NDSA-HR-{100+n} - الإصدار 2.0",
                    f"تاريخ السريان {eff.isoformat()} ({hijri_ish(eff)})",
                    "",
                    "أولاً: النطاق",
                    f"تسري هذه السياسة على جميع منسوبي {ORG_AR} وتشمل {desc_ar}.",
                    "وتُطبق بما لا يتعارض مع الأنظمة واللوائح ذات العلاقة.",
                    "",
                    "ثانياً: المرجعية النظامية",
                    f"تستند هذه السياسة إلى {ref_ar} الصادر عن الهيئة السعودية للبيانات والذكاء الاصطناعي.",
                    "ويُرجع إليه عند أي تعارض في التفسير.",
                    "",
                    "ثالثاً: الأحكام",
                    f"البند الأول: يكون الحد الأقصى للاستحقاق {pct} بالمئة من الراتب الأساسي.",
                    f"البند الثاني: تُقدَّم الطلبات قبل مدة لا تقل عن {days} يوم عمل.",
                    "البند الثالث: يعتمد الطلب مدير الإدارة المباشر ثم إدارة الموارد البشرية.",
                    "البند الرابع: لا يجوز الجمع بين هذا الاستحقاق وأي استحقاق مماثل عن المدة ذاتها.",
                    "",
                    "رابعاً: حماية البيانات الشخصية",
                    "تُعالج البيانات الشخصية المرتبطة بتطبيق هذه السياسة وفق مبدأ تقليل البيانات،",
                    "ولا يجوز مشاركتها خارج الإدارة المختصة إلا بموجب أساس نظامي.",
                    "وتُحفظ السجلات مدة 5 سنوات ثم تُتلف وفق دليل إتلاف البيانات الشخصية.",
                    "",
                    "خامساً: المسؤوليات",
                    "إدارة الموارد البشرية: تطبيق السياسة ومراجعتها سنوياً.",
                    "الإدارة القانونية: التحقق من التوافق النظامي.",
                    "",
                    DISC_AR,
                ]
                path = self.v1 / "hr" / f"{doc_id}.pdf"
                write_scanned_pdf(path, lines, seed=self.rng.randrange(10**6))
                transcript = "\n".join(l.lstrip("# ").strip() for l in lines if l.strip())
                (self.gold / f"{doc_id}.txt").parent.mkdir(parents=True, exist_ok=True)
                (self.gold / f"{doc_id}.txt").write_text(transcript, encoding="utf-8")
                self.record(doc_id=doc_id, family="hr_policy_scanned_ar",
                            path=str(path.relative_to(self.out)), format="pdf_scanned",
                            language="ar", department="hr", access_tier="restricted",
                            lifecycle="current", effective_date=eff.isoformat(),
                            title_ar=t_ar, title_en=t_en, cites=[ref_id],
                            ocr_gold=f"ocr_gold/{doc_id}.txt",
                            facts={"max_entitlement_pct": pct, "notice_days": days,
                                   "retention_years": 5})

    def hr_digital_en(self, count: int):
        made = 0
        for topic_id, t_en, t_ar, desc_en, desc_ar in HR_TOPICS:
            for variant in range(2):
                if made >= count:
                    return
                made += 1
                n = made
                doc_id = f"ndsa-hr-{topic_id}-en-{variant+1}"
                ref_id, ref_en, ref_ar = self.a_ref()
                eff = self.eff_date(30 * n + 7)
                pct = 15 + (n * 3) % 25
                days = 20 + (n * 2) % 15
                blocks = [
                    ("h", "1. Purpose and scope"),
                    ("p", f"This policy applies to all staff of {ORG_EN} and governs {desc_en}. "
                          f"It is issued under document number NDSA-HR-{200+n}, version 2.0, "
                          f"effective {eff.isoformat()}."),
                    ("h", "2. Regulatory basis"),
                    ("p", f"This policy implements the requirements of the <b>{ref_en}</b> issued by the "
                          f"Saudi Data and AI Authority. Where this policy and that instrument diverge, "
                          f"the instrument prevails."),
                    ("h", "3. Provisions"),
                    ("p", f"3.1 The maximum entitlement is <b>{pct}%</b> of basic salary."),
                    ("p", f"3.2 Requests shall be submitted no later than <b>{days} working days</b> "
                          f"before the intended start date."),
                    ("p", "3.3 Approval requires the line manager and then Human Resources. "
                          "Approval authority may not be sub-delegated below grade 11."),
                    ("p", "3.4 This entitlement may not be combined with any equivalent entitlement "
                          "for the same period."),
                    ("h", "4. Personal data"),
                    ("p", "Personal data processed under this policy is limited to what is necessary "
                          "for the stated purpose. Records are retained for <b>5 years</b> from the end "
                          "of the relevant financial year and then destroyed with evidence."),
                    ("h", "5. Responsibilities"),
                    ("table", [["Role", "Responsibility", "Review cycle"],
                               ["Human Resources", "Apply and maintain this policy", "Annual"],
                               ["Legal Affairs", "Confirm regulatory alignment", "Annual"],
                               ["Data Governance", "Classify records produced", "Semi-annual"]]),
                ]
                path = self.v1 / "hr" / f"{doc_id}.pdf"
                write_digital_pdf(path, f"{t_en}", f"{ORG_EN} · NDSA-HR-{200+n} · v2.0", blocks,
                                  DISC_EN)
                self.record(doc_id=doc_id, family="hr_policy_digital_en",
                            path=str(path.relative_to(self.out)), format="pdf_digital",
                            language="en", department="hr", access_tier="restricted",
                            lifecycle="current", effective_date=eff.isoformat(),
                            title_en=t_en, title_ar=t_ar, cites=[ref_id],
                            facts={"max_entitlement_pct": pct, "notice_days": days,
                                   "retention_years": 5, "min_approval_grade": 11})

    def allowance_tables(self, count: int):
        grades = ["Grade 8", "Grade 9", "Grade 10", "Grade 11", "Grade 12", "Grade 13", "Grade 14"]
        kinds = [
            ("housing", "Housing Allowance Bands", "بدل السكن", "SAR / month", 4500, 1800),
            ("transport", "Transport Allowance Bands", "بدل النقل", "SAR / month", 900, 260),
            ("perdiem-gcc", "Per-Diem — GCC destinations", "البدل اليومي - دول الخليج", "SAR / day", 750, 90),
            ("perdiem-intl", "Per-Diem — international destinations", "البدل اليومي - دولي", "SAR / day", 1150, 130),
            ("mobile", "Mobile and Connectivity Allowance", "بدل الاتصالات", "SAR / month", 300, 60),
            ("remote-setup", "Remote Work Setup Grant", "منحة تجهيز العمل عن بُعد", "SAR / one-off", 3000, 400),
            ("training", "Training Sponsorship Ceilings", "سقوف الابتعاث التدريبي", "SAR / year", 32000, 6000),
            ("relocation", "Relocation Support Bands", "بدل الانتقال", "SAR / one-off", 18000, 2500),
        ]
        for i, (kid, t_en, t_ar, unit, base, step) in enumerate(kinds[:count]):
            doc_id = f"ndsa-fin-allowance-{kid}"
            eff = self.eff_date(45 * (i + 1))
            rows, facts = [], {}
            for gi, g in enumerate(grades):
                amount = base + step * gi
                rows.append([g, f"{amount:,}", unit, f"{ORG_AR} - {t_ar}"])
                facts[f"{kid}_{g.lower().replace(' ', '')}"] = amount
            path = self.v1 / "finance" / f"{doc_id}.xlsx"
            write_xlsx(path, t_en, ["Grade", "Amount", "Unit", "Arabic label"], rows,
                       [f"Effective date: {eff.isoformat()}",
                        f"Regulatory basis: Data Classification Policy — this sheet is classified RESTRICTED.",
                        DISC_EN])
            self.record(doc_id=doc_id, family="allowance_tables",
                        path=str(path.relative_to(self.out)), format="xlsx",
                        language="bilingual", department="finance", access_tier="restricted",
                        lifecycle="current", effective_date=eff.isoformat(),
                        title_en=t_en, title_ar=t_ar, cites=["data-classification-policy"],
                        facts=facts)

    def it_procedures(self, count: int):
        made = 0
        for topic_id, t_en, desc in IT_TOPICS:
            for variant in range(2):
                if made >= count:
                    return
                made += 1
                n = made
                doc_id = f"ndsa-it-{topic_id}-{variant+1}"
                ref_id, ref_en, _ = self.a_ref()
                eff = self.eff_date(21 * n)
                hours = [4, 8, 12, 24, 48, 72][n % 6]
                sections = [
                    ("1. Purpose", [f"This procedure governs {desc} at {ORG_EN}.",
                                    f"Procedure reference NDSA-IT-{300+n}, revision {variant+1}.0, "
                                    f"effective {eff.isoformat()}."]),
                    ("2. Regulatory basis", [
                        f"This procedure operationalises the {ref_en} published by SDAIA.",
                        "Where a personal-data breach is involved, notification obligations under the "
                        "Personal Data Protection Law and its Implementing Regulation apply in addition "
                        "to the internal steps below."]),
                    ("3. Steps", [
                        "- Intake: the request or event is logged in the service management system with "
                        "a unique reference.",
                        f"- Triage: assessed within {hours} hours by the duty engineer against the "
                        "severity matrix in section 5.",
                        "- Approval: changes affecting classified data require Data Governance sign-off "
                        "before execution.",
                        "- Execution: performed by an authorised operator; a second operator verifies.",
                        "- Closure: evidence attached to the record; the record is retained per the "
                        "Logging and Monitoring Standard."]),
                    ("4. Prohibited actions", [
                        "- Sending data classified Secret or above to any external service, including "
                        "public generative AI tools.",
                        "- Exporting personal data outside the Kingdom without a completed transfer "
                        "risk assessment and documented approval.",
                        "- Using shared or generic accounts for any action in this procedure."]),
                    ("5. Severity and escalation", [
                        f"Severity 1 incidents escalate to the CIO immediately and to the Data "
                        f"Protection Officer within {hours} hours where personal data is implicated.",
                        "Severity 2 and 3 follow the standard service-level targets."]),
                    ("6. Records", [
                        "All records produced under this procedure are classified INTERNAL unless they "
                        "contain personal data, in which case they are classified RESTRICTED.",
                        "Retention is 3 years from closure."]),
                ]
                path = self.v1 / "it" / f"{doc_id}.docx"
                write_docx(path, t_en, sections, DISC_EN)
                self.record(doc_id=doc_id, family="it_procedures",
                            path=str(path.relative_to(self.out)), format="docx",
                            language="en", department="it", access_tier="internal",
                            lifecycle="current", effective_date=eff.isoformat(),
                            title_en=t_en, cites=[ref_id],
                            facts={"triage_hours": hours, "retention_years": 3})

    def intranet(self, count: int):
        made = 0
        for topic_id, t_en, t_ar in INTRANET_TOPICS:
            for variant in range(2):
                if made >= count:
                    return
                made += 1
                n = made
                doc_id = f"ndsa-web-{topic_id}-{variant+1}"
                ref_id, ref_en, ref_ar = self.a_ref()
                body_en = [
                    f"This page explains {t_en.lower()} at {ORG_EN}. It is a summary only; the "
                    f"authoritative text is the relevant policy document.",
                    f"Related regulation: {ref_en} (SDAIA).",
                    "If you cannot find what you need, contact the service desk on extension 4400.",
                    "Last reviewed by the Internal Communications team.",
                ]
                body_ar = [
                    f"توضح هذه الصفحة {t_ar} في {ORG_AR}. وهي ملخص فقط، والنص المعتمد هو وثيقة السياسة ذات العلاقة.",
                    f"المرجع النظامي: {ref_ar} (سدايا).",
                    "للاستفسار يرجى التواصل مع مكتب الخدمات على التحويلة ٤٤٠٠.",
                ]
                path = self.v1 / "intranet" / f"{doc_id}.html"
                write_html(path, t_en, t_ar, body_en, body_ar, crumb=t_en)
                self.record(doc_id=doc_id, family="intranet_pages",
                            path=str(path.relative_to(self.out)), format="html",
                            language="bilingual", department="general", access_tier="public",
                            lifecycle="current", effective_date=self.eff_date(10 * n).isoformat(),
                            title_en=t_en, title_ar=t_ar, cites=[ref_id],
                            facts={"extension": 4400})

    def circulars(self, count: int, pairs: int):
        """Identifier-rich circulars. The first `pairs` subjects get a v1
        (superseded, ships in corpus_v1) and a v2 (superseding, lands in
        corpus_v2 on Day 4) -- the whole F8 staleness drill lives here."""
        made = 0
        subj = CIRCULAR_SUBJECTS
        for i in range(count):
            s_en, s_ar = subj[i % len(subj)]
            year = 2024 + (i % 2)
            number = 10 + i
            circ_id = f"{number}/{year}"
            doc_id = f"ndsa-circ-{number}-{year}"
            ref_id, ref_en, ref_ar = self.a_ref()
            superseded = i < pairs
            eff_v1 = self.eff_date(14 * i)
            old_val = 20 + (i * 5) % 40
            new_val = old_val + 10

            def blocks_for(value, version, note):
                return [
                    ("h", f"Circular {circ_id} — {s_en}"),
                    ("p", f"<b>Arabic subject / الموضوع:</b> {s_ar}"),
                    ("p", f"To: all directorates. From: Office of the Governor, {ORG_EN}. "
                          f"Reference: NDSA-CIRC-{number}/{year}. Version: {version}."),
                    ("h", "Decision"),
                    ("p", f"With effect from the date below, the applicable figure is set at "
                          f"<b>{value}</b> (previously governed by the corresponding policy). "
                          f"All directorates shall apply this from the effective date without exception."),
                    ("p", note),
                    ("h", "Regulatory basis"),
                    ("p", f"Issued in alignment with the {ref_en} (SDAIA)."),
                    ("table", [["Field", "Value"],
                               ["Circular number", circ_id],
                               ["Version", version],
                               ["Applicable figure", str(value)],
                               ["Issuing office", "Office of the Governor"]]),
                ]

            if superseded:
                note_v1 = ("This circular is SUPERSEDED by a later circular of the same number. "
                           "Do not apply it to decisions taken after the superseding date.")
                p1 = self.v1 / "circulars" / f"{doc_id}-v1.pdf"
                write_digital_pdf(p1, f"Circular {circ_id} (v1) — {s_en}",
                                  f"{ORG_EN} · effective {eff_v1.isoformat()} · SUPERSEDED",
                                  blocks_for(old_val, "1.0", note_v1), DISC_EN)
                self.record(doc_id=f"{doc_id}-v1", family="circulars",
                            path=str(p1.relative_to(self.out)), format="pdf_digital",
                            language="bilingual", department="legal", access_tier="internal",
                            lifecycle="current",
                            lifecycle_by_corpus={"v1": "current", "v2": "superseded"},
                            effective_date=eff_v1.isoformat(),
                            superseded_by=f"{doc_id}-v2", circular_id=circ_id,
                            title_en=f"Circular {circ_id} — {s_en}", title_ar=s_ar,
                            cites=[ref_id], corpus=["v1", "v2"],
                            facts={"applicable_figure": old_val})

                eff_v2 = eff_v1 + timedelta(days=400)
                note_v2 = (f"This circular supersedes version 1.0 of circular {circ_id} in its "
                           f"entirety. The previous figure of {old_val} no longer applies.")
                p2 = self.v2 / "circulars" / f"{doc_id}-v2.pdf"
                write_digital_pdf(p2, f"Circular {circ_id} (v2) — {s_en}",
                                  f"{ORG_EN} · effective {eff_v2.isoformat()} · CURRENT",
                                  blocks_for(new_val, "2.0", note_v2), DISC_EN)
                self.record(doc_id=f"{doc_id}-v2", family="circulars",
                            path=str(p2.relative_to(self.out)), format="pdf_digital",
                            language="bilingual", department="legal", access_tier="internal",
                            lifecycle="current",
                            lifecycle_by_corpus={"v2": "current"},
                            effective_date=eff_v2.isoformat(),
                            supersedes=f"{doc_id}-v1", circular_id=circ_id,
                            title_en=f"Circular {circ_id} — {s_en}", title_ar=s_ar,
                            cites=[ref_id], corpus=["v2"],
                            facts={"applicable_figure": new_val})
            else:
                p = self.v1 / "circulars" / f"{doc_id}.pdf"
                write_digital_pdf(p, f"Circular {circ_id} — {s_en}",
                                  f"{ORG_EN} · effective {eff_v1.isoformat()} · CURRENT",
                                  blocks_for(old_val, "1.0", "This circular remains in force."),
                                  DISC_EN)
                self.record(doc_id=doc_id, family="circulars",
                            path=str(p.relative_to(self.out)), format="pdf_digital",
                            language="bilingual", department="legal", access_tier="internal",
                            lifecycle="current", effective_date=eff_v1.isoformat(),
                            circular_id=circ_id, title_en=f"Circular {circ_id} — {s_en}",
                            title_ar=s_ar, cites=[ref_id], corpus=["v1", "v2"],
                            facts={"applicable_figure": old_val})
            made += 1

    # -- driver ------------------------------------------------------------
    def build(self):
        fam = {f["id"]: f for f in self.spec["families"]}
        self.hr_scanned_ar(fam["hr_policy_scanned_ar"]["count"])
        self.hr_digital_en(fam["hr_policy_digital_en"]["count"])
        self.allowance_tables(fam["allowance_tables"]["count"])
        self.it_procedures(fam["it_procedures"]["count"])
        self.intranet(fam["intranet_pages"]["count"])
        c = fam["circulars"]
        self.circulars(c["count"], c["supersession"]["pairs"])

        idx = {
            "spec_version": self.spec["spec_version"],
            "seed": self.spec["seed"],
            "organisation": self.spec["organisation"],
            "generated_documents": len(self.index),
            "documents": self.index,
        }
        (self.out / "layer_b_index.json").write_text(
            json.dumps(idx, ensure_ascii=False, indent=2), encoding="utf-8")
        return idx


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--spec", type=Path, default=ROOT / "corpus/layer_b_specs/layer_b.yaml")
    ap.add_argument("--out", type=Path, default=ROOT / "data/corpus")
    args = ap.parse_args()

    spec = yaml.safe_load(args.spec.read_text(encoding="utf-8"))
    b = Builder(spec, args.out)
    idx = b.build()

    by_fmt: dict[str, int] = {}
    for d in idx["documents"]:
        by_fmt[d["format"]] = by_fmt.get(d["format"], 0) + 1
    print(f"Layer B: {idx['generated_documents']} documents -> {args.out}")
    for k, v in sorted(by_fmt.items()):
        print(f"  {k:14s} {v}")
    sup = sum(1 for d in idx["documents"]
              if (d.get("lifecycle_by_corpus") or {}).get("v2") == "superseded")
    print(f"  superseded pairs: {sup}")
    print(f"  ocr_gold transcripts: {len(list((args.out / 'ocr_gold').glob('*.txt')))}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
