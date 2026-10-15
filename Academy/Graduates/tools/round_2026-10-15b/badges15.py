# -*- coding: utf-8 -*-
# Round 15 Oct 2026 (b) — SDAIA Badges: tag badge holders (badge name + level), enrich programmes, add 1 new record
import json,re,os,sys,shutil,collections,copy
import openpyxl
DRY='--apply' not in sys.argv
H=os.path.expanduser('~/mnt/Academy/'); B='backup-2026-10-15b'
DATE='15 أكتوبر 2026'; TAG='[شارات سدايا — 15 أكتوبر 2026] '
SRC='كشف الحاصلين على شارات سدايا (الدفعة الأولى — اختبار 25 يوليو 2026) — ملف داخلي من الأكاديمية، 15 أكتوبر 2026'
COH='الدفعة الأولى – 25 يوليو 2026'
LV={'خبير':('Expert','expert'),'مشارك':('Associate','associate'),'مساعد':('Assistant','assistant')}
TRK={'مهندس الذكاء الاصطناعي':'AI Engineer','عالم البيانات':'Data Scientist'}
def badge(track,lvl):
    return {"ar":f"{track} – {lvl}","en":f"{TRK[track]} – {LV[lvl][0]}","lvl":LV[lvl][1],"tr":'ai' if TRK[track]=='AI Engineer' else 'ds',"date":"2026-07-25","cohort":COH}
AIE='مهندس الذكاء الاصطناعي'
P_AML='أساليب تعلم الآلة المتقدمة — أكاديمية سدايا'
P_MDE='هندسة البيانات الحديثة لأنظمة الذكاء الاصطناعي — أكاديمية سدايا'
P_BDA='بناء وكلاء البيانات (From Data to Dialogue) — أكاديمية سدايا'
P_AWF='القوى العاملة المدعومة بالوكلاء: تحويل تجارب الذكاء الاصطناعي الأولية إلى إنتاجية الأعمال — أكاديمية سدايا'
P_ANLP='معسكر معالجة اللغة العربية — أكاديمية سدايا'
P_TS='تحليل السلاسل الزمنية والتنبؤ — أكاديمية سدايا'
P_AAG='هندسة أنظمة الذكاء الاصطناعي التوكيلي المتقدمة — أكاديمية سدايا'
P_LLS='بناء النماذج اللغوية الكبيرة من الصفر — أكاديمية سدايا'
P_NVL='برنامج التدريب الاحترافي للنماذج اللغوية الكبيرة (إنفيديا) — أكاديمية سدايا'
# (sheet name, record id or None, level, programmes listed in the badge sheet, programmes to APPEND to prog, match note)
ROWS=[
 ('يزيد محمد علي العبيدان',198,'خبير',['برنامج التدريب الاحترافي في الذكاء الاصطناعي التوليدي (NVIDIA)'],[],'مطابقة تامة للاسم الرباعي'),
 ('تهاني محمد سبيت السبيت',292,'مشارك',['أساليب تعلم الآلة المتقدمة'],[P_AML],'الاسم الأول والأب والعائلة'),
 ('سارة يوسف محمد الشريف',None,'مشارك',['برنامج التدريب الاحترافي للنماذج اللغوية الكبيرة (إنفيديا)'],[],'سجل جديد — لا مطابق'),
 ('لينا خليل إبراهيم الضويحي',389,'مشارك',['التدريب الاحترافي في النماذج اللغوية الكبيرة'],[],'مطابقة تامة للاسم الرباعي'),
 ('هتون سعد عبدالرحمن الشعيل',227,'مشارك',['برنامج التدريب الاحترافي للنماذج اللغوية الكبيرة (إنفيديا)'],[],'الاسم الأول والعائلة + مسار NVIDIA'),
 ('عبدالرحمن فهد ناصر بن هليل',228,'مساعد',['هندسة البيانات الحديثة لأنظمة الذكاء الاصطناعي','برنامج التدريب الاحترافي للنماذج اللغوية الكبيرة (إنفيديا)'],[P_MDE],'الاسم الأول والعائلة + مسار NVIDIA'),
 ('ابتهال سالم صالح باعظيم',295,'مساعد',['بناء وكلاء البيانات','القوى العاملة المدعومة بالوكلاء'],[P_BDA,P_AWF],'الاسم الأول والأب والعائلة (عائلة نادرة)'),
 ('عبدالعزيز حسن عبدالله ملا',380,'مساعد',['التدريب الاحترافي في النماذج اللغوية الكبيرة'],[],'الاسم الرباعي (حسن/حسان اختلاف إملائي) + مسار NVIDIA'),
 ('هياء بعيان فايز الوزره',390,'مساعد',['معسكر معالجة اللغة العربية','التدريب الاحترافي في النماذج اللغوية الكبيرة'],[P_ANLP],'الاسم الرباعي؛ #1793 «هياء الوزره» على الأرجح الشخص نفسه'),
 ('محمد على سعيد الحارثي',241,'مساعد',['برنامج التدريب الاحترافي للنماذج اللغوية الكبيرة (إنفيديا)'],[],'الاسم الأول والعائلة + مسار NVIDIA (مرشح وحيد) — ثقة متوسطة'),
 ('شهد ناصر حمدان الهاجري',13,'مساعد',['تحليل السلاسل الزمنية والتنبؤ','هندسة أنظمة الذكاء الاصطناعي التوكيلي المتقدمة','التدريب الاحترافي في النماذج اللغوية الكبيرة'],[P_TS,P_AAG],'الاسم الأول والعائلة + مسار NVIDIA'),
 ('ميعاد سلطان سعد الشيباني',376,'مساعد',['هندسة البيانات الحديثة لأنظمة الذكاء الاصطناعي','بناء النماذج اللغوية الكبيرة من الصفر','التدريب الاحترافي في النماذج اللغوية الكبيرة'],[P_MDE,P_LLS],'مطابقة تامة للاسم الرباعي'),
 ('رغد محمد عبدالله العواجي',348,'مساعد',['هندسة أنظمة الذكاء الاصطناعي التوكيلي المتقدمة'],[P_AAG],'الاسم الأول والأب والعائلة'),
 ('ضحى أحمد حمد العبدالوهاب',3223,'مساعد',['برنامج التدريب الاحترافي للنماذج اللغوية الكبيرة (إنفيديا)'],[P_NVL],'الاسم الأول والعائلة (اسم نادر)'),
]
FILES=['SuccessStories/website/js/data.js','index.html','Graduates/index.html','Graduates/Graduates_Database.xlsx','Graduates/Graduates_Database.md']
if not DRY:
    for f in FILES:
        a,b=os.path.splitext(H+f); t=f'{a}.{B}{b}'
        if not os.path.exists(t): shutil.copy2(H+f,t)
def rep(s,a,b,count=None):
    n=s.count(a); assert n>=1,('missing',a[:90])
    if count is not None: assert n==count,(a[:90],n)
    return s.replace(a,b)
P=H+'Graduates/index.html'; s=open(P,encoding='utf-8').read()
i=s.index('const DATA = ')+len('const DATA = '); j=s.index('\nconst SITE_STATS')
D=json.loads(s[i:j].rstrip()[:-1]); g=D['grads']; N0=len(g); assert N0==3375,N0
assert not any('badge' in r for r in g)
head=s[:i]; tail=s[j:]
T0=collections.Counter(r['tr'] for r in g)
X_=H+'Graduates/Graduates_Database.xlsx'; wb=openpyxl.load_workbook(X_); ws=wb['قاعدة البيانات']
assert ws.cell(N0+1,1).value==N0 and ws.max_column==23
# new column 24
hc=ws.cell(1,24); hc.value='شارة سدايا — SDAIA Badge'; src=ws.cell(1,23)
hc.font=copy.copy(src.font); hc.fill=copy.copy(src.fill); hc.alignment=copy.copy(src.alignment); hc.border=copy.copy(src.border)
ws.column_dimensions['X'].width=30
catl={'program':'برنامج','project':'مشروع','cert':'شهادة','job':'توظيف'}
log=[]
for nm,k,lvl,progs,app,mnote in ROWS:
    bd=badge(AIE,lvl)
    txt=(f"حاصل/ة على شارة سدايا «{bd['ar']}» ({bd['en']}) — {COH}. البرامج المسجّلة في كشف الشارات: "+' · '.join(progs)+
         f". الاسم في الكشف: «{nm}». المطابقة: {mnote}.")
    if k is None:
        k=len(g)+1
        rec={"n":nm,"en":"Sarah Yousef Mohammed Alsharif","lv":"badge1015","score":80,"prog":P_NVL,"co":"قبل يوليو 2026 (اختبار الشارة: 25 يوليو 2026)",
             "win":"ضمن النطاق","edu":"","emp":"","role":"","cats":["program","cert"],
             "d":TAG+txt+' تنويه: سجل #243 «ساره الحارثي الشريف» (برامج NVIDIA) يشترك في الاسم الأول والعائلة لكن اسم الأب مختلف (يوسف/الحارثي) — إشارة تقاطع فقط، لا دمج. الاسم الإنجليزي نقل حرفي عن الاسم العربي الرسمي — غير متحقق علنًا.',
             "links":[],"src":SRC,"tr":"genai","badge":bd}
        g.append(rec); r=k+1
        vals=[k,rec['n'],rec['en'],rec['prog'],'أكاديمية الذكاء الاصطناعي التوليدي — Gen AI Academy',rec['co'],None,None,None,None,rec['lv'],rec['score'],rec['win'],'برنامج · شهادة',None,None,None,None,None,None,rec['d'],None,SRC,bd['ar']]
        for ci,v in enumerate(vals,1): ws.cell(r,ci).value=v
        log.append((k,nm,rec['n'],bd,'سجل جديد',progs,mnote)); continue
    G=g[k-1]; r=k+1; assert ws.cell(r,1).value==k and ws.cell(r,2).value==G['n'],(k,ws.cell(r,2).value)
    G['badge']=bd; ws.cell(r,24).value=bd['ar']
    for p in app:
        core=p.split(' — ')[0]
        if core not in G['prog']: G['prog']+=' + '+p
    ws.cell(r,4).value=G['prog']
    if 'cert' not in G['cats']:
        G['cats'].append('cert'); cur=ws.cell(r,14).value; ws.cell(r,14).value=(cur+' · ' if cur else '')+'شهادة'
    G['d']=(G['d']+' ' if G['d'] else '')+TAG+txt; ws.cell(r,21).value=((ws.cell(r,21).value or '')+' '+TAG+txt).strip()
    log.append((k,nm,G['n'],bd,'إثراء + وسم',progs,mnote))
# cross-note on the likely duplicate of Haya
G=g[1793-1]; t=TAG+'على الأرجح الشخص نفسه في السجل #390 «هياء بعيان بن فايز الوزره» (كشف الشارات يذكر معسكر معالجة اللغة العربية لها) — وُسمت الشارة على #390 وحده؛ يُقترح دمج السجلين في جولة تنظيف لاحقة.'
G['d']=(G['d']+' ' if G['d'] else '')+t; ws.cell(1794,21).value=((ws.cell(1794,21).value or '')+' '+t).strip()
N=len(g); T=collections.Counter(r['tr'] for r in g); NB=sum(1 for r in g if r.get('badge'))
print('N',N0,'->',N,{k:(T0[k],T[k]) for k in T if T0[k]!=T[k]},'badges',NB)
for x in log: print(x[0],x[2],'|',x[3]['ar'],'|',x[4])
if DRY: sys.exit(0)
TOT=format(N,','); OTOT=format(N0,',')
# --- portal UI: badge tag, filter option, search haystack
tail=rep(tail,"${g.coop && g.tr !== 'coop' ? `<span class=\"sd-tag sd-tag--solid\" style=\"background:${TRACK_TAG.coop}\">${I('handshake', 'i-sm')}${T('التدريب التعاوني', 'Cooperative Training')}</span>` : ''}",
 "${g.coop && g.tr !== 'coop' ? `<span class=\"sd-tag sd-tag--solid\" style=\"background:${TRACK_TAG.coop}\">${I('handshake', 'i-sm')}${T('التدريب التعاوني', 'Cooperative Training')}</span>` : ''}\n    ${g.badge ? `<span class=\"sd-tag sd-tag--solid badge-tag\" style=\"background:#7A5300\" title=\"${T('شارة سدايا — ' + g.badge.cohort, 'SDAIA Badge — first cohort, 25 July 2026')}\">${I('badge-check', 'i-sm')}${T('شارة سدايا: ' + g.badge.ar, 'SDAIA Badge: ' + g.badge.en)}</span>` : ''}",1)
tail=rep(tail,"if (state.cat && !(g.cats || []).includes(state.cat)) return false;","if (state.cat && !(state.cat === 'badge' ? g.badge : (g.cats || []).includes(state.cat))) return false;",1)
tail=rep(tail,"+ ' ' + (g.role || '')).toLowerCase()","+ ' ' + (g.role || '') + (g.badge ? ' شارة سدايا sdaia badge ' + g.badge.ar + ' ' + g.badge.en : '')).toLowerCase()",1)
head=rep(head,'<option value="job" data-ar="توظيف" data-en="Employment"></option>','<option value="badge" data-ar="حاملو شارات سدايا" data-en="SDAIA Badge holders"></option>\n            <option value="job" data-ar="توظيف" data-en="Employment"></option>',1)
head=head.replace(OTOT,TOT); tail=tail.replace(OTOT,TOT)
tail=rep(tail,"{ value: %d, label: 'سجلًا فرديًا',"%N0,"{ value: %d, label: 'سجلًا فرديًا',"%N,1)
EN={'ai':'AI Engineer','ds':'Data Scientist','dmg':'Data Management & Governance','other':'Other / Specialized','genai':'Gen AI Academy'}
for k,en in EN.items():
    if T[k]!=T0[k]:
        tail,n=re.subn(r"(en: '%s', value: )%d\b"%(re.escape(en),T0[k]),"\\g<1>%d"%T[k],tail); print('tail',k,n); assert n>=1
D['grads']=g; open(P,'w',encoding='utf-8').write(head+json.dumps(D,ensure_ascii=False)+';'+tail)
P2=H+'index.html'; s=open(P2,encoding='utf-8').read().replace(OTOT,TOT)
s=rep(s,'data-count="%d"'%N0,'data-count="%d"'%N,1)
for k in ('ai','ds','dmg','other','genai','coop'):
    pc='%.1f'%(100*T[k]/N)
    pat=r'(data-count=")%d(">0</div>.*?--w:)[\d.]+(%%.*?data-ar=")[\d.]+(%% من السجل" data-en=")[\d.]+(%% of the registry")'%T0[k]
    s,n=re.subn(pat,lambda mm:mm.group(1)+str(T[k])+mm.group(2)+pc+mm.group(3)+pc+mm.group(4)+pc+mm.group(5),s,count=1,flags=re.S); assert n==1,k
open(P2,'w',encoding='utf-8').write(s)
P3=H+'SuccessStories/website/js/data.js'; s=open(P3,encoding='utf-8').read()
s=rep(s,"{ value: %d, label: 'خريجًا موثّقًا',"%N0,"{ value: %d, label: 'خريجًا موثّقًا',"%N,1)
for k,en in EN.items():
    if T[k]!=T0[k]:
        s,n=re.subn(r"(en: '%s',\s+value: )%d\b"%(re.escape(en),T0[k]),"\\g<1>%d"%T[k],s); assert n==1,(k,n)
s=s.replace(OTOT,TOT)
open(P3,'w',encoding='utf-8').write(s)
# --- xlsx stats + round sheet
wl=sum(1 for r in g if r.get('links')); uu=len({l[1].rstrip('/') for r in g for l in (r.get('links') or [])})
hi=sum(1 for r in g if r.get('score',0)>=75); mid=sum(1 for r in g if 45<=r.get('score',0)<75); lo=N-hi-mid
emp=sum(1 for r in g if r.get('emp')); yc=collections.Counter()
for r in g:
    ys=set(re.findall(r'20(2[1-6])',r.get('co') or ''))
    if not ys: yc['none']+=1
    for y in ys: yc['20'+y]+=1
st=wb['إحصاءات']
M={'الإجمالي':TOT,'مهندس ذكاء اصطناعي':T['ai'],'عالم بيانات':T['ds'],'أكاديمية الذكاء الاصطناعي التوليدي':T['genai'],'إدارة وحوكمة البيانات':T['dmg'],'التدريب التعاوني':T['coop'],'برامج متخصصة أخرى':T['other'],
   'ثقة مرتفعة (≥75)':hi,'ثقة متوسطة (45–74)':mid,'ثقة منخفضة (<45)':lo,'سجلات بجهة عمل موثّقة':emp,'سجلات بروابط تحقق علنية':wl,'روابط مصدر فريدة':uu,
   'سجلات بدفعة 2021':yc['2021'],'سجلات بدفعة 2022':yc['2022'],'سجلات بدفعة 2023':yc['2023'],'سجلات بدفعة 2024':yc['2024'],'سجلات بدفعة 2025':yc['2025'],'سجلات بدفعة 2026':yc['2026'],'سجلات بلا دفعة محددة':yc['none'],
   'تحديث':f'{DATE} (ب) — وسم حاملي شارات سدايا: {NB} سجلًا موسومًا (13 قائمًا + 1 جديد)؛ {OTOT} → {TOT}'}
seen=set()
for row in st.iter_rows():
    if row[0].value in M: row[1].value=M[row[0].value]; seen.add(row[0].value)
bl='حاملو شارات سدايا'
found=[row for row in st.iter_rows() if row[0].value==bl]
if found: found[0][1].value=NB
else: st.append([bl,NB])
sh=wb.create_sheet('شارات_سدايا_15_أكتوبر_2026')
sh.append(['#','الاسم في كشف الشارات','الاسم في القاعدة','الشارة','Badge','المستوى','دفعة الاختبار','البرامج (من الكشف)','الإجراء','المطابقة'])
for k,nm,dn,bd,act,progs,mn in log:
    sh.append([k,nm,dn,bd['ar'],bd['en'],bd['ar'].split(' – ')[1],COH,' · '.join(progs),act,mn])
sh.sheet_view.rightToLeft=True
wb.save(X_)
# --- md
P5=H+'Graduates/Graduates_Database.md'; s=open(P5,encoding='utf-8').read()
s=rep(s,'**الإصدار:** 15 أكتوبر 2026 — ','**الإصدار:** 15 أكتوبر 2026 (ب) — وسم حاملي شارات سدايا (%s → %s · %d سجلًا موسومًا). قبله: 15 أكتوبر 2026 — '%(OTOT,TOT,NB),1)
L=['## وسم حاملي شارات سدايا — %s (ب)'%DATE,'',f'**الخريجون {OTOT} → {TOT} · {NB} سجلًا يحمل وسم «شارة سدايا» · القصص وجهات العمل دون تغيير.**','',
 'المصدر: كشف الأكاديمية للحاصلين على الشارات (14 اسمًا، الدفعة الأولى — اختبار 25 يوليو 2026). الشارات نوعان (مهندس الذكاء الاصطناعي · عالم البيانات) بثلاثة مستويات (مساعد · مشارك · خبير)؛ الكشف الحالي كله «مهندس الذكاء الاصطناعي»: خبير 1 · مشارك 4 · مساعد 9. الوسم حقل جديد «badge» في البوابة (شارة ذهبية على البطاقة + خيار «حاملو شارات سدايا» في قائمة التصفية + قابل للبحث) وعمود جديد (24) «شارة سدايا — SDAIA Badge» في ملف Excel. أرقام الهوية في الكشف لم تُنقل إلى القاعدة.','',
 '| # | الاسم في الكشف | الاسم في القاعدة | الشارة | الإجراء | المطابقة |','|---|---|---|---|---|---|']
for k,nm,dn,bd,act,progs,mn in log: L.append(f'| {k} | {nm} | {dn} | {bd["ar"]} | {act} | {mn} |')
L+=['','**ملاحظات:** أُضيفت إلى حقل البرنامج البرامجُ الواردة في الكشف غير المسجّلة سابقًا (#292 · #228 · #295 · #390 · #13 · #376 · #348 · #3223)؛ وأُضيفت فئة «شهادة» حيث غابت. #1793 «هياء الوزره» على الأرجح تكرار لـ #390 — وُسم #390 وحده ويُقترح الدمج لاحقًا. #3376 سارة يوسف محمد الشريف سجل جديد (مسار أكاديمية الذكاء الاصطناعي التوليدي، كبقية سجلات برنامج NVIDIA)، مع إشارة تقاطع إلى #243 «ساره الحارثي الشريف» دون دمج لاختلاف اسم الأب. #241 «محمد الحارثي» مطابقة بثقة متوسطة (اسم ثنائي شائع، المرشح الوحيد في مسار NVIDIA).','','---','']
k=s.index('## جولة الخريجين الجدد وفحص القصص والتدقيق والتنظيف — 15 أكتوبر 2026')
s=s[:k]+'\n'.join(L)+'\n'+s[k:]
open(P5,'w',encoding='utf-8').write(s)
for f in ('index.html','Graduates/index.html','SuccessStories/website/js/data.js'):
    t=open(H+f,encoding='utf-8').read(); print(f,'leftover',OTOT,t.count(OTOT),'new',t.count(TOT))
print('done',N,NB)
