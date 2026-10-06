# Round 6 Oct 2026 — deep sweep (LinkedIn people search -> experience pages) + story check + final cleaning round
import json,re,os,sys,shutil,collections,difflib
sys.path.insert(0,os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0,os.path.expanduser('~/mnt/Academy/Graduates/tools'))
import load, verdicts as V, newrecs as NR, translit as TR
import openpyxl
DRY='--apply' not in sys.argv
H=os.path.expanduser('~/mnt/Academy/')
B='backup-2026-10-06'
DATE='6 أكتوبر 2026'
SRC='بحث الأشخاص على LinkedIn (حساب مسجَّل) + صفحات الخبرة — 6 أكتوبر 2026'
TAG='[جولة المسح العميق — 6 أكتوبر 2026] '
CTAG='[جولة التنظيف — 6 أكتوبر 2026] '
TL=' تنويه: الاسم العربي نقل حرفي عن الاسم اللاتيني المنشور — غير متحقق علنًا.'
VERIFIED_AR={'lama-al-zakri-%D9%84%D9%80%D9%85%D9%89-%D8%A7%D9%84%D8%B2%D9%83%D8%B1%D9%8A-502087200'}
TRN={'ai':'مهندس ذكاء اصطناعي — AI Engineer','ds':'عالم بيانات — Data Scientist','other':'برامج متخصصة أخرى — Other / Specialized','dmg':'إدارة وحوكمة البيانات — Data Management & Governance'}
LI=lambda s:'https://www.linkedin.com/in/'+s+'/'
def rep(s,a,b,count=None):
    n=s.count(a); assert n>=1,('missing',a[:90])
    if count is not None: assert n==count,(a[:90],n)
    return s.replace(a,b)
def rrep(s,pat,b,count=1,flags=0):
    s2,n=re.subn(pat,b,s,flags=flags); assert n==count,(pat,n); return s2
def bk(p):
    a,b=os.path.splitext(p); t=f'{a}.{B}{b}'
    if not os.path.exists(t): shutil.copy2(p,t)
FILES=['SuccessStories/website/js/data.js','SuccessStories/website/index.html','index.html','Graduates/index.html','Graduates/Graduates_Database.xlsx','Graduates/Graduates_Database.md']
if not DRY:
    for f in FILES: bk(H+f)

# ---------------- load
P=H+'Graduates/index.html'; s=open(P,encoding='utf-8').read()
i=s.index('const DATA = ')+len('const DATA = '); j=s.index('\nconst SITE_STATS')
raw=s[i:j].rstrip(); D=json.loads(raw[:-1]); g=D['grads']; N0=len(g); assert N0==3223,N0
head=s[:i]; tail=s[j:]
X=load.stories(); ST=X['S']; assert len(ST)==125
X_=H+'Graduates/Graduates_Database.xlsx'
wb=openpyxl.load_workbook(X_); ws=wb['قاعدة البيانات']
assert ws.cell(N0+1,1).value==N0
def wsd(i_,txt):
    r=i_+1; assert ws.cell(r,1).value==i_; ws.cell(r,21).value=(ws.cell(r,21).value or '')+txt

def slug(u):
    m=re.search(r'linkedin\.com/in/([^/?#]+)',u); return m.group(1).lower().rstrip('/') if m else None
SL={}
for k,r in enumerate(g,1):
    for l in r['links']:
        sl=slug(l[1])
        if sl: SL[sl]=k

# ---------------- dedupe new records
idx=TR.build_index(g)
def norm_en(x): return re.sub(r'[^a-z]','',x.lower().replace('al-','al').replace('al ','al'))
ENM=collections.defaultdict(list)
for k,r in enumerate(g,1):
    if r.get('en'): ENM[norm_en(r['en'])].append(k)
dupes=[]; NEWL=[]
for rec in NR.R:
    en,ar,sl=rec[0],rec[1],rec[2]
    hit=set()
    if sl.lower() in SL: hit.add(('slug',SL[sl.lower()]))
    for k in ENM.get(norm_en(en),[]): hit.add(('en',k))
    try:
        for k in (TR.match(ar,en,idx) or []): hit.add(('tr',k if isinstance(k,int) else k[0]))
    except Exception as e: print('tr err',e)
    if hit: dupes.append((rec,sorted(hit)))
    NEWL.append((rec,hit))
print('candidate duplicates:',len(dupes))
for rec,h in dupes:
    print(' ',rec[0],'|',rec[1],'|',rec[3],rec[4],'=>',[(t,k,g[k-1]['n'],g[k-1]['en'],g[k-1]['prog'][:40],g[k-1]['co'][:20],[l[1][-30:] for l in g[k-1]['links']][:1]) for t,k in h])
# story map check (names)
SN={norm_en(x.get('nameEn','')) for x in ST}; SA={x['name'] for x in ST}
for rec,_ in NEWL:
    if rec[6]=='PASS' and (norm_en(rec[0]) in SN or rec[1] in SA): print('ALREADY STORY?',rec[0],rec[1])
for i_,p in V.PASS.items():
    G=g[i_-1]
    if norm_en(G['en']) in SN or G['n'] in SA: print('ALREADY STORY (existing rec)?',i_,G['n'])
if DRY:
    json.dump([[r[0],sorted(h)] for r,h in NEWL if h],open('dupes.json','w'),ensure_ascii=False)
    sys.exit(0)

# =============== APPLY ===============
# merges decided after dry-run review: rec en-name -> existing id (enrichment instead of new record)
MERGE=json.load(open('merge_decisions.json')) if os.path.exists('merge_decisions.json') else {}
GEN_F_FIRST={'Maali','Noura','Jana','Mona','Dhuha','Ghada','Suad','Layan','Khuzama','Nada','Raghad','Abeer','Mjd','Renad','Shouq','Fatimah','Amal','Leen','Radhyah','Maria','Faten','Awatef','Anwar'}
def fem(rec): return rec[0].split()[0] in GEN_F_FIRST
PSHORT={'T5':'معسكر T5','DSB':'معسكر T5 لعلوم البيانات','ZEHAM':'معسكر تقنيات إدارة الزحام','DM':'معسكر إدارة البيانات','DG':'معسكر حوكمة البيانات','LLM':'معسكر النماذج اللغوية الكبيرة','ML':'معسكر تعلم الآلة','MLAI':'معسكر تعلم الآلة','AISC':'برنامج أبطال صيف الذكاء الاصطناعي','FG':'معسكر حديثي التخرج','SNOW':'معسكر ServiceNow','OXF':'معسكر أكسفورد','AIPRO':'معسكر محترفي الذكاء الاصطناعي','APPAI':'معسكر الذكاء الاصطناعي التطبيقي','AIENG':'معسكر مهندس الذكاء الاصطناعي','AIAPPS':'برنامج بناء تطبيقات الذكاء الاصطناعي','ADVDL':'برنامج المسارات المتقدمة','T5ALLAM':'معسكر T5 وبرنامج علّام','T5ALLAM2':'معسكر T5 وبرنامج علّام','FGAISC':'معسكر حديثي التخرج','MLAISC':'معسكر تعلم الآلة','LLMAISC':'معسكر النماذج اللغوية الكبيرة'}
def kebab(en):
    return re.sub(r'[^a-z0-9]+','-',en.lower().replace('al-','al').replace('.','')).strip('-')
IDS={x['id'] for x in ST}
def uid(base):
    b=base; n=2
    while b in IDS: b=f'{base}-{n}'; n+=1
    IDS.add(b); return b
def role_ar(role): return role.split(' (')[0]
def gapword(gp):
    m=re.match(r'\+(\d+)',gp); n=int(m.group(1)) if m else None
    W={1:'شهر واحد',2:'شهرين',3:'ثلاثة أشهر',4:'أربعة أشهر',5:'خمسة أشهر',6:'ستة أشهر',7:'سبعة أشهر',8:'ثمانية أشهر',9:'تسعة أشهر',10:'عشرة أشهر',11:'أحد عشر شهرًا',12:'اثني عشر شهرًا',13:'ثلاثة عشر شهرًا',14:'أربعة عشر شهرًا'}
    return W.get(n,gp)
stories=[]; new_ids=[]; enrich=[]
for rec,hit in NEWL:
    en,ar,sl,pk,co,entry,verdict,stv=rec
    prog,tr=NR.P[pk]
    vtxt=verdict if verdict!='PASS' else f'PASS_DATED ({stv[4]}) — {stv[2]} — {stv[0]} ({stv[3]}).'
    if verdict=='CERT_ONLY': vtxt='CERT_ONLY — لا نتيجة لاحقة مؤرخة في صفحة الخبرة.'
    if verdict=='NO_COMPLETION_DATE': vtxt='NO_COMPLETION_DATE — قيد المعسكر بلا تاريخ ظاهر.'
    if en in MERGE:
        k=int(MERGE[en]); G=g[k-1]
        t=' '+TAG+f'إثراء: صفحة الخبرة على LinkedIn ({LI(sl)}): {entry}. [فحص القصص] {vtxt}'
        G['d']+=t; wsd(k,t)
        if not any(slug(l[1])==sl.lower() for l in G['links']):
            G['links'].append(['LinkedIn',LI(sl)]); r=k+1; cur=ws.cell(r,15).value; ws.cell(r,15).value=(cur+' · ' if cur else '')+LI(sl)
        if not G.get('en'): G['en']=en; ws.cell(k+1,3).value=en
        rid=k; enrich.append(k)
    else:
        rid=len(g)+1; r=rid+1
        d=f'صفحة الخبرة على LinkedIn (حساب مسجَّل، {DATE}): {entry}. عتبة الهوية مستوفاة بسمتين: الاسم على الملف الشخصي + قيد تدريب مؤرخ يسمّي برنامج الأكاديمية تحت جهة SDAIA. [فحص القصص] {vtxt}'
        if hit: d+=' إشارة متبادلة غير مدموجة: '+' · '.join(sorted({f'#{k}' for _,k in hit}))+' (اسم مشابه بلا رابط مشترك — لا دمج).'
        if sl not in VERIFIED_AR and re.search('[؀-ۿ]',ar): d+=TL
        score=84 if sl in VERIFIED_AR else 80
        G={"n":ar,"en":en,"lv":"li1006","score":score,"prog":prog,"co":co,"win":"ضمن النطاق","edu":"","emp":"","role":"","cats":["program"],"d":d,"links":[["LinkedIn",LI(sl)]],"src":SRC,"tr":tr}
        g.append(G); new_ids.append(rid)
        vals=[rid,ar,en,prog,TRN[tr],co,None,None,None,None,'li1006',score,'ضمن النطاق','برنامج',LI(sl),None,None,None,None,None,d,None,SRC]
        for ci,v in enumerate(vals,1): ws.cell(r,ci).value=v
    if verdict=='PASS' and en=='Mjd Alotaibi':
        t=f' {TAG}قصة «majd-alotaibi» قائمة؛ صفحة الخبرة لا تُظهر «تحكّم» (المصدر الوسيط المحذوف)، فأُعيد تثبيت القصة على جاهز (يونيو 2025، +8) وصُحّح الدليل.'; G['d']+=t; wsd(rid,t); G['emp']='جاهز (Jahez)'; ws.cell(rid+1,7).value='جاهز (Jahez)'
    elif verdict=='PASS':
        org,logo,role,hire,gp,year,period,extra=stv
        G['emp']=org; G['role']=role
        if 'job' not in G['cats']: G['cats'].append('job')
        ws.cell(rid+1,7).value=org; ws.cell(rid+1,8).value=role
        f=fem(rec); t_='ت' if f else ''
        first=G['n'].split()[0] if re.search('[؀-ۿ]',G['n']) else en.split()[0]
        ra=role_ar(role)
        undis=org.startswith('جهة حكومية')
        impact=f"التحاق{'ها' if f else 'ه'} بـ{org} «{ra}» في {hire} — بعد {gapword(gp)} من الإتمام." if not undis else f"التحاق{'ها' if f else 'ه'} بجهة حكومية «{ra}» في {hire} — بعد {gapword(gp)} من الإتمام."
        story=f"أتمّ{t_} {first} {prog} ({period}). وفي {hire} — بعد {gapword(gp)} من الإتمام — التحق{t_} {'بجهة حكومية' if undis else 'بـ'+org} {ra}{'ة' if False else ''} بدوام كامل." + (f" {extra}." if extra else '') + " النتيجة المحتسبة هي هذا التعيين، المؤرخ بصفحة الخبرة على LinkedIn."
        story=story.replace('..','.')
        q=f"ما تعلّمته في {PSHORT[pk]} بأكاديمية سدايا كان طريقي إلى العمل {'' if True else ''}{ra.replace('أخصائية','أخصائيةً').replace('محللة','محللةً').replace('مهندسة','مهندسةً').replace('عالمة','عالمةً').replace('مطوّرة','مطوّرةً').replace('مستشارة','مستشارةً')}."
        sid=uid(kebab(en))
        st={"id":sid,"year":year,"quote":{"text":q,"kind":"draft"},"name":G['n'],"nameEn":en,"photo":None,
            "role":role,"org":org,"orgLogo":logo,"program":prog,"period":period,"category":"employment","categories":["employment"],
            "impact":impact,"story":story,"achievements":[prog,f"{role} — {org} ({hire})"],
            "links":[{"label":"LinkedIn","url":LI(sl)}],
            "provenance":f"صفحة الخبرة على LinkedIn بحساب مسجَّل (قُرئت {DATE}): {entry} · {role} — {org} ({hire}) + سجل الأكاديمية #{rid}."}
        stories.append((rid,st))
        t=f' {TAG}مُرقّى إلى قصة نجاح «{sid}».'; G['d']+=t; wsd(rid,t) if rid<=N0 else None
        if rid>N0: ws.cell(rid+1,21).value+=t

# ---------------- existing records: story-shortlist verdicts
for i_,p in V.PASS.items():
    G=g[i_-1]; r=i_+1
    if i_ in V.CO:
        co,_=V.CO[i_]; t=f' {TAG}تصويب الدفعة من صفحة الخبرة/الشهادات (كانت: «{G["co"]}»).'; G['co']=co; ws.cell(r,6).value=co; G['d']+=t; wsd(i_,t)
    if not any(slug(l[1])==p['slug'].lower() for l in G['links']):
        G['links'].append(['LinkedIn',LI(p['slug'])]); cur=ws.cell(r,15).value; ws.cell(r,15).value=(cur+' · ' if cur else '')+LI(p['slug'])
    t=f" {TAG}صفحة الخبرة (حساب مسجَّل): {p['exp']}. PASS_DATED ({p['gap']}) — مُرقّى إلى قصة نجاح «{p['id']}»."
    G['d']+=t; wsd(i_,t); G['emp']=p['org']; G['role']=p['role']; ws.cell(r,7).value=p['org']; ws.cell(r,8).value=p['role']
    if 'job' not in G['cats']: G['cats'].append('job')
    st={"id":uid(p['id']),"year":p['year'],"quote":{"text":p['quote'],"kind":"draft"},"name":G['n'],"nameEn":G['en'],"photo":None,
        "role":p['role'],"org":p['org'],"orgLogo":p['logo'],"program":p['program'],"period":p['period'],"category":"employment","categories":["employment"],
        "impact":p['impact'],"story":p['story'],"achievements":[p['program'],f"{p['role']} — {p['org']} ({p['hire']})"],
        "links":[{"label":"LinkedIn","url":LI(p['slug'])}],
        "provenance":f"صفحة الخبرة على LinkedIn بحساب مسجَّل (قُرئت {DATE}): {p['exp']} + سجل الأكاديمية #{i_}."}
    stories.append((i_,st))
for i_,v in V.HOLD.items():
    t=f' {TAG}HOLD — قيد التحقق: {v}'; g[i_-1]['d']+=t; wsd(i_,t)
for i_,v in V.FAIL.items():
    t=f' {TAG}صفحة الخبرة (حساب مسجَّل): {v}'; g[i_-1]['d']+=t; wsd(i_,t)
for i_ in V.CERT:
    t=f' {TAG}صفحة الخبرة (حساب مسجَّل): CERT_ONLY — لا يظهر بعد البرنامج إلا قيد التدريب نفسه أو تدريب/تدريب تعاوني، بلا تعيين مؤرخ ذي صلة.'; g[i_-1]['d']+=t; wsd(i_,t)

# ---------------- CLEANING
clean=collections.Counter()
# (a) dead LinkedIn links: remove from links, keep URL in d
for i_ in V.DEAD:
    G=g[i_-1]; dead=[l for l in G['links'] if 'linkedin.com/in/' in l[1]]
    G['links']=[l for l in G['links'] if 'linkedin.com/in/' not in l[1]]
    t=f' {CTAG}رابط LinkedIn لم يعد متاحًا (صفحة 404 في {DATE}) فأُزيل من الروابط: '+' · '.join(l[1] for l in dead)+'.'
    G['d']+=t; wsd(i_,t); ws.cell(i_+1,15).value=None; clean['dead_link']+=1
# shatha #2119: slug changed
G=g[2119-1]
if 'غامدي' in G['n']:
    G['links'].append(['LinkedIn',LI('shatha-al-ghamdi-66808a337')]); ws.cell(2120,15).value=LI('shatha-al-ghamdi-66808a337')
    t=f' {CTAG}الرابط المُحدَّث: shatha-al-ghamdi-66808a337 (المعرّف الرقمي نفسه 66808a337).'; G['d']+=t; wsd(2119,t); clean['slug_updated']+=1
# (b) country-subdomain + http normalization
for k,G in enumerate(g,1):
    ch=False
    for l in G['links']:
        u=l[1]; u2=re.sub(r'^https?://[a-z]{2}\.linkedin\.com/','https://www.linkedin.com/',u)
        if u2.startswith('http://'): u2='https://'+u2[7:]
        if u2!=u: l[1]=u2; ch=True; clean['url_norm']+=1
    if ch and k<=N0:
        for c in (15,16,17,18,19,20):
            v=ws.cell(k+1,c).value
            if isinstance(v,str):
                v2=re.sub(r'https?://[a-z]{2}\.linkedin\.com/','https://www.linkedin.com/',v).replace('http://scholar','https://scholar')
                if v2!=v: ws.cell(k+1,c).value=v2
# (c) placeholder employers -> blank
PH={'غير موثّق علنًا','—','غير ظاهرة علنًا','غير موثّق علنًا (دفعة حديثة)'}
for k,G in enumerate(g,1):
    if G['emp'] in PH:
        G['emp']=''; clean['emp_placeholder']+=1
        if k<=N0: ws.cell(k+1,7).value=None
    if G['emp']=='SDAIA ¦ سدايا':
        G['emp']='SDAIA | سدايا'; clean['emp_variant']+=1
        if k<=N0: ws.cell(k+1,7).value='SDAIA | سدايا'
# (d) cross-notes for same-name groups not yet cross-referenced
def normar(t):
    t=re.sub('[إأآا]','ا',t); t=t.replace('ة','ه').replace('ى','ي'); t=re.sub(r'[ً-ْ]','',t); return ' '.join(t.lower().split())
nm=collections.defaultdict(list)
for k,G in enumerate(g,1): nm[normar(G['n'])].append(k)
for key,ids in nm.items():
    if len(ids)<2: continue
    for k in ids:
        others=[o for o in ids if o!=k and f'#{o}' not in g[k-1]['d']]
        if others:
            t=f' {CTAG}إشارة متبادلة: الاسم نفسه في '+' · '.join(f'#{o}' for o in others)+' — برنامج أو دفعة مختلفة ولا رابط شخصي مشترك؛ لا دمج (البند 7).'
            g[k-1]['d']+=t
            if k<=N0: wsd(k,t)
            else: ws.cell(k+1,21).value+=t
            clean['xnote']+=1
# (e) hygiene asserts
for G in g:
    assert '**' not in G['d']
    G['d']=re.sub(r'  +',' ',G['d'])
print('clean',dict(clean))

# ---------------- stats
N=len(g); T=collections.Counter(r['tr'] for r in g)
wl=sum(1 for r in g if r.get('links')); uu=len({l[1].rstrip('/') for r in g for l in (r.get('links') or [])})
emp=sum(1 for r in g if r.get('emp'))
hi=sum(1 for r in g if r.get('score',0)>=75); mid=sum(1 for r in g if 45<=r.get('score',0)<75); lo=N-hi-mid
yc=collections.Counter()
for r in g:
    ys=set(re.findall(r'20(2[1-6])',r.get('co') or ''))
    if not ys: yc['none']+=1
    for y in ys: yc['20'+y]+=1
EXCL_RE=re.compile(r'غير معلنة|^جهة حكومية|قمة الابتكار|OSACT7|Private Data Department|فريق|مركز .*مسابقة')
def emps(S):
    return {x['org'] for x in S if not EXCL_RE.search(x['org'])}
E0=len(emps(ST)); print('E0 computed',E0)
ALL=ST+[x for _,x in stories]; E1c=len(emps(ALL))
E1=75+(E1c-E0)
S0=125; S1=S0+len(stories)
print('N',N,dict(T),'wl',wl,'uu',uu,'emp',emp,'S1',S1,'E1',E1,'new',len(new_ids),'enrich',len(enrich),'years',dict(yc))
TOT=format(N,',')
OLD=dict(N=3223,TOT='3,223',S=125,E=75,ai=1560,ds=885,dmg=328,other=190,wl=1274,uu='2,159')
VER_AR_OLD='إصدار 3 أكتوبر 2026 (جولة مسح منشورات LinkedIn وفحص القصص'
VER_AR_NEW='إصدار 6 أكتوبر 2026 (جولة المسح العميق والتنظيف النهائي'
VER_EN_OLD='version 3 October 2026 (LinkedIn post sweep and story-check round'
VER_EN_NEW='version 6 October 2026 (deep sweep and final cleaning round'
def common(t):
    t=t.replace(OLD['TOT'],TOT).replace(VER_AR_OLD,VER_AR_NEW).replace(VER_EN_OLD,VER_EN_NEW)
    t=t.replace(f"({OLD['S']} قصة)",f'({S1} قصة)').replace(f"({OLD['S']} stories)",f'({S1} stories)').replace(f"{OLD['S']} قصة نجاح موثّقة",f'{S1} قصة نجاح موثّقة').replace(f"{OLD['S']} verified success stories",f'{S1} verified success stories')
    t=t.replace(f"عبر {OLD['E']} جهة عمل",f'عبر {E1} جهة عمل').replace(f"across {OLD['E']} employers",f'across {E1} employers')
    return t
head=common(head); tail=common(tail)
tail=rep(tail,"{ value: %d, label: 'سجلًا فرديًا',"%OLD['N'],"{ value: %d, label: 'سجلًا فرديًا',"%N,1)
for k,en in (('ai','AI Engineer'),('ds','Data Scientist'),('dmg','Data Management & Governance'),('other','Other / Specialized')):
    tail=rrep(tail,r"(en: '%s', value: )%d\b"%(re.escape(en),OLD[k]),"\\g<1>%d"%T[k])
tail=rrep(tail,r"\{ value: %d,(\s+)label: 'سجلًا برابط تحقق علني', note: '%s رابط مصدر فريد',"%(OLD['wl'],OLD['uu']),"{ value: %d,\\1label: 'سجلًا برابط تحقق علني', note: '%s رابط مصدر فريد',"%(wl,format(uu,',')))
tail=rep(tail,"noteEn: '%s unique source links'"%OLD['uu'],"noteEn: '%s unique source links'"%format(uu,','),1)
tail=rrep(tail,r"\{ value: %d,(\s+)label: 'قصة نجاح موثّقة'"%OLD['S'],"{ value: %d,\\1label: 'قصة نجاح موثّقة'"%S1)
D['grads']=g
open(P,'w',encoding='utf-8').write(head+json.dumps(D,ensure_ascii=False)+';'+tail)
print('grad leftovers',{k:(head+tail).count(k) for k in [OLD['TOT'],'3223',OLD['uu'],f"{OLD['S']} قصة",VER_AR_OLD]})
# portal
P2=H+'index.html'; s=open(P2,encoding='utf-8').read()
s=common(s)
s=rep(s,'data-count="%d"'%OLD['N'],'data-count="%d"'%N,1)
s=rep(s,'data-count="%d"'%OLD['S'],'data-count="%d"'%S1,1); s=rep(s,'data-count="%d"'%OLD['E'],'data-count="%d"'%E1,1)
s=s.replace(f"{OLD['S']} قصة",f'{S1} قصة').replace(f"{OLD['S']} stories",f'{S1} stories').replace(f"{OLD['E']} جهة عمل",f'{E1} جهة عمل').replace(f"{OLD['E']} employers",f'{E1} employers').replace(f"across {OLD['E']} government",f'across {E1} government')
for k,old in (('ai',OLD['ai']),('ds',OLD['ds']),('dmg',OLD['dmg']),('other',OLD['other']),('genai',249),('coop',11)):
    pc='%.1f'%(100*T[k]/N)
    pat=r'(data-count=")%d(">0</div>.*?--w:)[\d.]+(%%.*?data-ar=")[\d.]+(%% من السجل" data-en=")[\d.]+(%% of the registry")'%old
    s,n=re.subn(pat,lambda m:m.group(1)+str(T[k])+m.group(2)+pc+m.group(3)+pc+m.group(4)+pc+m.group(5),s,count=1,flags=re.S); assert n==1,k
open(P2,'w',encoding='utf-8').write(s)
print('portal leftovers',{k:s.count(k) for k in [OLD['TOT'],str(OLD['N']),f"{OLD['S']} قصة",f"{OLD['E']} جهة"]})
# data.js
P3=H+'SuccessStories/website/js/data.js'; s=open(P3,encoding='utf-8').read()
s=rep(s,"{ value: %d, label: 'خريجًا موثّقًا',"%OLD['N'],"{ value: %d, label: 'خريجًا موثّقًا',"%N,1)
for k,en in (('ai','AI Engineer'),('ds','Data Scientist'),('dmg','Data Management & Governance'),('other','Other / Specialized')):
    s=rrep(s,r"(en: '%s',\s+value: )%d\b"%(re.escape(en),OLD[k]),"\\g<1>%d"%T[k])
s=rrep(s,r"(\{ value: )%d(,\s+label: 'قصة نجاح)"%OLD['S'],r"\g<1>%d\2"%S1)
s=rrep(s,r"(\{ value: )%d(,\s+label: 'جهة عمل')"%OLD['E'],r"\g<1>%d\2"%E1)
k=s.index('const STORIES = ['); e=s.index('\n];',k)
s=s[:e]+'\n'+',\n'.join('  '+json.dumps(x,ensure_ascii=False) for _,x in stories)+','+s[e:]
s=rep(s,"'use strict';","/* %s: جولة المسح العميق والتنظيف النهائي — 3,223 ← %s خريجًا (%d سجلًا جديدًا · %d إثراءً)؛ %d قصة جديدة (125 ← %d)، وجهات العمل 75 ← %d. */\n'use strict';"%(DATE,TOT,len(new_ids),len(enrich),len(stories),S1,E1),1)
MAJD={"id":"majd-alotaibi","year":2024,"quote":{"text":"معسكر الذكاء الاصطناعي وعلوم البيانات في أكاديمية سدايا فتح لي الطريق إلى العمل عالمَ بيانات في جاهز.","kind":"draft"},
 "name":"مجد العتيبي","nameEn":"Majd Alotaibi","photo":"assets/majd-alotaibi.jpg","role":"عالم بيانات (Data Scientist)","org":"جاهز (Jahez)","orgLogo":"assets/jahez.png",
 "program":"معسكر الذكاء الاصطناعي وعلوم البيانات (T5) — أكاديمية سدايا","period":"يوليو – أكتوبر 2024","category":"employment","categories":["employment"],
 "impact":"التحاق بجاهز «عالم بيانات» بدوام كامل في يونيو 2025 — بعد ثمانية أشهر من إتمام المعسكر.",
 "story":"أتمّ مجد معسكر الذكاء الاصطناعي وعلوم البيانات في أكاديمية سدايا (يوليو–أكتوبر 2024)، ثم واصل التطوير بمعسكرَي الذكاء الاصطناعي التوليدي وبناء وتطوير نماذج الذكاء الاصطناعي في أكاديمية طويق (يناير–مايو 2025). وفي يونيو 2025 — بعد ثمانية أشهر من الإتمام — التحق بشركة جاهز عالمَ بيانات بدوام كامل. النتيجة المحتسبة هي التعيين في جاهز، المؤرخ بصفحة الخبرة على LinkedIn.",
 "achievements":["معسكر الذكاء الاصطناعي وعلوم البيانات (T5) — أكاديمية سدايا (يوليو–أكتوبر 2024)","عالم بيانات (Data Scientist) — جاهز (يونيو 2025)","معسكرا الذكاء الاصطناعي التوليدي وبناء النماذج — أكاديمية طويق (2025)"],
 "links":[{"label":"LinkedIn","url":"https://www.linkedin.com/in/majd-abdullah-/"},{"label":"GitHub","url":"https://github.com/majdalotaibi"}],
 "provenance":"تصويب 6 أكتوبر 2026: كانت القصة مثبتة على «عالم بيانات في تحكّم 2024» من مصدر وسيط (أُزيل في 27 أغسطس)؛ صفحة الخبرة على LinkedIn (حساب مسجَّل) لا تُظهر تحكّم، وتُظهر: Artificial Intelligence & Data Science Bootcamp — SDAIA (يوليو–أكتوبر 2024) · Tuwaiq Academy (يناير–مايو 2025) · Data Scientist — Jahez International Company (Full-time، منذ يونيو 2025). أُعيد تثبيت القصة على جاهز (+8). سجل الأكاديمية #2332."}
a=s.index("  {\n    id: 'majd-alotaibi'"); b=s.index("\n  },\n",a)+len("\n  },\n")
s=s[:a]+'  '+json.dumps(MAJD,ensure_ascii=False)+',\n'+s[b:]
assert s.count("majd-alotaibi'")==0
open(P3,'w',encoding='utf-8').write(s)
P4=H+'SuccessStories/website/index.html'; s=open(P4,encoding='utf-8').read(); s=rep(s,'"numberOfItems": %d'%OLD['S'],'"numberOfItems": %d'%S1,1); open(P4,'w',encoding='utf-8').write(s)
# xlsx stats + sheet
st=wb['إحصاءات']
M={'الإجمالي':TOT,'مهندس ذكاء اصطناعي':T['ai'],'عالم بيانات':T['ds'],'أكاديمية الذكاء الاصطناعي التوليدي':T['genai'],'إدارة وحوكمة البيانات':T['dmg'],'التدريب التعاوني':T['coop'],'برامج متخصصة أخرى':T['other'],
   'ثقة مرتفعة (≥75)':hi,'ثقة متوسطة (45–74)':mid,'ثقة منخفضة (<45)':lo,'سجلات بجهة عمل موثّقة':emp,'سجلات بروابط تحقق علنية':wl,'روابط مصدر فريدة':uu,
   'سجلات بدفعة 2021':yc['2021'],'سجلات بدفعة 2022':yc['2022'],'سجلات بدفعة 2023':yc['2023'],'سجلات بدفعة 2024':yc['2024'],'سجلات بدفعة 2025':yc['2025'],'سجلات بدفعة 2026':yc['2026'],'سجلات بلا دفعة محددة':yc['none'],
   'قصص النجاح':S1,'جهات العمل':E1,
   'تحديث':f'{DATE} — جولة المسح العميق والتنظيف النهائي: {len(new_ids)} سجلًا جديدًا (3,223 → {TOT})، {len(enrich)} إثراءً؛ {len(stories)} قصة نجاح جديدة (125 → {S1})، وجهات العمل 75 → {E1}'}
for row in st.iter_rows():
    if row[0].value in M: row[1].value=M[row[0].value]
sh=wb.create_sheet('جولة_المسح_العميق_2026-10-06')
sh.append(['#','الاسم','النوع','البرنامج / الحكم','الدليل'])
for nid in new_ids: G=g[nid-1]; sh.append([nid,G['n'],'سجل جديد',G['prog'],G['links'][0][1]])
for k in enrich: sh.append([k,g[k-1]['n'],'إثراء','',''])
for rid,x in stories: sh.append([rid,x['name'],'قصة جديدة',x['org']+' — '+x['role'],x['links'][0]['url']])
for i_,v in V.HOLD.items(): sh.append([i_,g[i_-1]['n'],'قيد التحقق',v,''])
for i_,v in V.FAIL.items(): sh.append([i_,g[i_-1]['n'],'مرفوض',v.split(' — ')[0],''])
for i_ in V.DEAD: sh.append([i_,g[i_-1]['n'],'تنظيف: رابط ميت أُزيل','',''])
for k_,v in clean.items(): sh.append(['','',f'تنظيف: {k_}',v,''])
wb.save(X_)
json.dump(dict(N=N,T=T,wl=wl,uu=uu,new=new_ids,enrich=enrich,stories=[[r,x['id'],x['org']] for r,x in stories],S1=S1,E1=E1,clean=clean,yc=yc,emp=emp),open('stats.json','w'),ensure_ascii=False)
# md
P5=H+'Graduates/Graduates_Database.md'; s=open(P5,encoding='utf-8').read()
s=rep(s,'**الإصدار:** 3 أكتوبر 2026 — ','**الإصدار:** %s — جولة المسح العميق والتنظيف النهائي (3,223 → %s · القصص 125 → %d · جهات العمل 75 → %d). قبله: 3 أكتوبر 2026 — '%(DATE,TOT,S1,E1),1)
L=['## جولة المسح العميق والتنظيف النهائي — %s'%DATE,'',f'**الخريجون 3,223 → {TOT} · القصص 125 → {S1} · جهات العمل 75 → {E1}.**','',
   f'**1) سجلات جديدة ({len(new_ids)})** — بحث الأشخاص على LinkedIn (18 استعلامًا) ثم صفحة الخبرة: قيد تدريب مؤرخ يسمّي برنامج الأكاديمية تحت جهة SDAIA. استُبعد من لا يظهر اسمه إلا بالأحرف الأولى، ومن قيده تدريب عام في سدايا (تعاوني/متدرب) دون برنامج للأكاديمية.','',
   '| # | الاسم | الاسم الإنجليزي | البرنامج | الدفعة | الرابط |','|---|---|---|---|---|---|']
for nid in new_ids:
    G=g[nid-1]; L.append(f"| {nid} | {G['n']} | {G['en']} | {G['prog']} | {G['co']} | {G['links'][0][1]} |")
L+=['',f'**2) إثراءات ({len(enrich)}):** '+' · '.join('#%d'%k for k in enrich)+'.','',f'**3) قصص نجاح جديدة ({len(stories)}):**','','| القصة | # | النتيجة |','|---|---|---|']
for rid,x in stories: L.append(f"| {x['id']} | {rid} | {x['role']} — {x['org']} |")
L+=['',f'**قيد التحقق من قائمة الفحص السابقة ({len(V.HOLD)}):** '+' · '.join('#%d'%k for k in V.HOLD)+'.','',
    f'**4) التنظيف النهائي:** '+' · '.join(f'{k}: {v}' for k,v in clean.items())+'. الروابط الميتة: '+' · '.join('#%d'%k for k in V.DEAD)+'.','','---','']
k=s.index('## جولة مسح منشورات LinkedIn وفحص القصص — 3 أكتوبر 2026')
s=s[:k]+'\n'.join(L)+'\n'+s[k:]
open(P5,'w',encoding='utf-8').write(s)
print('done')
