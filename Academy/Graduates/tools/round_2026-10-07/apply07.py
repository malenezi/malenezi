# -*- coding: utf-8 -*-
# Round 7 Oct 2026 — gap-closing round (stories + DB fields + GitHub open item)
import json,re,os,sys,shutil,collections
sys.path.insert(0,os.path.dirname(os.path.abspath(__file__)))
import load, fills as F
import openpyxl
DRY='--apply' not in sys.argv
H=os.path.expanduser('~/mnt/Academy/')
B='backup-2026-10-07'
DATE='7 أكتوبر 2026'
TAG='[جولة سد الفجوات — 7 أكتوبر 2026] '
TL=' تنويه: الاسم العربي نقل حرفي عن الاسم اللاتيني المنشور — غير متحقق علنًا.'
SRC='مسح مستودعات GitHub (بند مفتوح من جولة 6 أكتوبر) + بحث LinkedIn بحساب مسجَّل — 7 أكتوبر 2026'
TRN={'ai':'مهندس ذكاء اصطناعي — AI Engineer','ds':'عالم بيانات — Data Scientist','other':'برامج متخصصة أخرى — Other / Specialized','dmg':'إدارة وحوكمة البيانات — Data Management & Governance'}
def rep(s,a,b,count=None):
    n=s.count(a); assert n>=1,('missing',a[:90])
    if count is not None: assert n==count,(a[:90],n)
    return s.replace(a,b)
def rrep(s,pat,b,count=1,flags=0):
    s2,n=re.subn(pat,b,s,flags=flags); assert n==count,(pat,n); return s2
def bk(p):
    a,b=os.path.splitext(p); t=f'{a}.{B}{b}'
    if not os.path.exists(t): shutil.copy2(p,t)
FILES=['SuccessStories/website/js/data.js','SuccessStories/website/index.html','index.html','Graduates/index.html','Graduates/Graduates_Database.xlsx','Graduates/Graduates_Database.md']
if not DRY:
    for f in FILES: bk(H+f)
P=H+'Graduates/index.html'; s=open(P,encoding='utf-8').read()
i=s.index('const DATA = ')+len('const DATA = '); j=s.index('\nconst SITE_STATS')
raw=s[i:j].rstrip(); D=json.loads(raw[:-1]); g=D['grads']; N0=len(g); assert N0==3309,N0
head=s[:i]; tail=s[j:]
X=load.stories(); ST=X['S']; assert len(ST)==168
X_=H+'Graduates/Graduates_Database.xlsx'
wb=openpyxl.load_workbook(X_); ws=wb['قاعدة البيانات']
assert ws.cell(N0+1,1).value==N0
def wsd(k,txt):
    r=k+1; assert ws.cell(r,1).value==k; ws.cell(r,21).value=(ws.cell(r,21).value or '')+txt
def note(k,t):
    g[k-1]['d']+=t; wsd(k,t)
def addlink(k,label,url,col):
    G=g[k-1]; norm=lambda u:re.sub(r'^https?://(www\.|[a-z]{2}\.)?','',u).rstrip('/').lower()
    if any(norm(l[1])==norm(url) for l in G['links']): return False
    G['links'].append([label,url]); r=k+1; cur=ws.cell(r,col).value; ws.cell(r,col).value=(cur+' · ' if cur else '')+url; return True
C=collections.Counter(); LOG=collections.defaultdict(list)
def yr(c): return re.findall(r'20(2[1-6])',c or '')
# ---------- 1) DB: cohort/window fills
for k,(co,ev) in F.CO.items():
    G=g[k-1]; r=k+1; old=G['co']
    if co and co!=old:
        G['co']=co; ws.cell(r,6).value=co; C['co_set']+=1
        t=f' {TAG}تأريخ الدفعة من LinkedIn (حساب مسجَّل): {ev} — حُدّثت الدفعة (كانت: «{old or "فارغة"}»).'
    else:
        t=f' {TAG}تأكيد التأريخ من LinkedIn (حساب مسجَّل): {ev}.'; C['co_confirmed']+=1
    note(k,t); LOG['co'].append(k)
for k,v in F.STAFF.items():
    note(k,f' {TAG}⚠ صفحة الخبرة (حساب مسجَّل): {v}، ولا يظهر قيد معسكر/برنامج مؤرخ للأكاديمية — البرنامج مستنتج؛ مرشح للاستبعاد بقاعدة «الموظفون والمدربون ليسوا خريجين» عند المراجعة.'); C['staff_flag']+=1; LOG['staff'].append(k)
for k,v in F.NOEV.items():
    note(k,f' {TAG}⚠ صفحة الخبرة والشهادات (حساب مسجَّل): {v} — لا دليل على الانتساب للأكاديمية؛ مرشح للاستبعاد عند المراجعة.'); C['noev_flag']+=1; LOG['staff'].append(k)
for k,G in enumerate(g,1):
    if G['win']=='غير محدد':
        ys=yr(G['co'])
        if ys and 'بدء موثّق' not in G['co'] and all(y in('21','22','24','25','26') for y in ys):
            G['win']='ضمن النطاق'; ws.cell(k+1,13).value='ضمن النطاق'; C['win_set']+=1; LOG['win'].append(k)
# ---------- 2) DB: English names
for k,en in F.EN.items():
    G=g[k-1]; assert not G['en'],(k,G['en'])
    G['en']=en; ws.cell(k+1,3).value=en; C['en_set']+=1
    note(k,f' {TAG}الاسم الإنجليزي من {F.EN_SRC[k]}: «{en}».')
# ---------- 3) DB: enrichments
ENR=[(9,None,None,'دُمجت قصة «amal-almubarak» في قصة «amal-almuarik» (الشخص نفسه — السجلان دُمجا في 5 سبتمبر لكن القصتين بقيتا)؛ إشارة متبادلة: #1493 «Amal Bint Abdulrahman Almuarik» (برنامج بناء تطبيقات الذكاء الاصطناعي، أكتوبر–نوفمبر 2024) — الاسم نفسه بلا رابط مشترك؛ لا دمج.'),
     (802,'GitHub','https://github.com/etab12','حساب GitHub «etab12» (Etab — AI | ML | Bioinformatics) فيه مستودع «SDAIA-LLM-Bootcamp» (فبراير 2024) — يطابق شهادة معسكر LLM (5 فبراير 2024).'),
     (1504,'GitHub','https://github.com/Mohammadalm','حساب GitHub «Mohammad Almubaddal» (Data scientist، الرياض) فيه «Pandas homework sdaia bootcamp» (يناير 2024) — اللقب نادر ويطابق الاسم الكامل في السجل.'),
     (3284,'GitHub','https://github.com/SarahAljuwayr','حساب GitHub «SarahAljuwayr» فيه مستودع «T5_Bootcamp» (يوليو 2024) — يطابق الدفعة.'),
     (3276,'GitHub','https://github.com/ebtisamasiri','حساب GitHub «Ebtsam» (ebtisamasiri) فيه «T5-Data-Science-AI-Practice: Projects, datasets, and code from my T5 Bootcamp» (أغسطس 2024) — يطابق الدفعة.'),
     (2647,'LinkedIn','https://www.linkedin.com/in/rundalsaleh/','ملف LinkedIn «Rund Alsaleh» (rundalsaleh): Digital Transformation Analyst — Sadara Chemical Company · Full-time منذ فبراير 2025، وشهادات NVIDIA (Transformer-Based NLP · Conversational AI · Deep Learning) — يطابق جهة العمل والمسمى في قائمة الأكاديمية (رند = Rund).'),
     (2643,'LinkedIn','https://www.linkedin.com/in/lamalz/','صفحة الخبرة (lamalz): Data Engineering Bootcamp — SDAIA (نوفمبر 2024 – أبريل 2025) ثم Marketing Analyst — stc · Full-time منذ سبتمبر 2025 (+5) — يطابق قائمة الأكاديمية.'),
     (2645,'LinkedIn','https://www.linkedin.com/in/aljohara-alhogail-08841425a/','صفحة الخبرة: Cloud Systems Management Bootcamp — SDAIA (أكتوبر–ديسمبر 2025، تكريم ICAN 2025) ثم RPA Developer — Alinma · Full-time منذ يناير 2026 (+1). تسمّي قائمة الأكاديمية «برنامج وكلاء الذكاء الاصطناعي».'),
     (2646,'LinkedIn','https://www.linkedin.com/in/badr-a-alshamrani-49a04a245/','صفحة الخبرة: Applied AI & Innovation Engineer — KFSHRC · Full-time منذ ديسمبر 2025 (+3) — مؤكد.'),
     (2642,'LinkedIn','https://www.linkedin.com/in/rawabi-abod-allihyni/','صفحة الخبرة: SADAIA & Microsoft Data Engineering Bootcamp (نوفمبر 2024 – أبريل 2025) ثم «محلل بيانات» — أمانة العاصمة المقدسة بصيغة Apprenticeship منذ مايو 2026 (+13).'),
     (3303,'GitHub','https://github.com/Manaralbogami','حساب GitHub «Manaralbogami» فيه مستودع «diabetes-prediction» بوصف «Data science Bootcamp SDAIA T5» (يناير 2022) — يطابق قيد المعسكر في صفحة الخبرة (ديسمبر 2021 – يناير 2022). [بند GitHub المفتوح من جولة 6 أكتوبر: تبيّن أنه سجل قائم لا سجل جديد.]'),
     (980,'GitHub','https://github.com/mo-100','حساب GitHub «Mohammed Alageel» (الرياض؛ Building production ML & NLP systems) فيه «SDAIA-ML-Bootcamp» (يناير 2024) و«T5-Data-Science-Bootcamp-projects» (أغسطس 2024) — يطابق معسكر تعلم الآلة (الربع الأول 2024) في السجل. ملف LinkedIn محتمل «mohammed-o-alageel» (Data Scientist — MOZN؛ O = عمر) لم يُنسب لغياب قيد سدايا الظاهر. [بند GitHub المفتوح: سجل قائم لا جديد.]'),
     (141,None,None,'التأريخ الشهري (LinkedIn): شهادتا NVIDIA-Certified Associate يونيو 2025 → AI Engineer — Arabot · Full-time منذ ديسمبر 2025 (+6).'),
     (308,None,None,'التأريخ الشهري (LinkedIn): شهادتا NVIDIA-Certified Associate فبراير 2025 → Agentic AI Engineer — IBM · Full-time (أكتوبر 2025 – يناير 2026، +8) → AI Engineer — مصرف الراجحي منذ فبراير 2026 (+12).'),
     (365,None,None,'التأريخ الشهري (LinkedIn): AI Specialist — Confidential Government · Full-time منذ يناير 2026 (+6 من آخر دورة NVIDIA، يوليو 2025).'),
     (802,None,None,'التأريخ الشهري (LinkedIn): شهادتا NVIDIA فبراير 2025 → AI Specialist — KAIMRC · Full-time منذ أبريل 2026 (+14، على حد النافذة).'),
     (2715,None,None,'صفحة الخبرة: مدير مشروع — وزارة الاتصالات وتقنية المعلومات · Full-time منذ مايو 2025 (+7) أُضيف إنجازًا في القصة؛ تبقى الجائزة النتيجة المحتسبة.'),
]
for k,lab,url,txt in ENR:
    if url:
        col=15 if lab=='LinkedIn' else 16
        if addlink(k,lab,url,col): C['link_added']+=1
    note(k,' '+TAG+txt); LOG['enrich'].append(k)
G=g[2647-1]
G['alias']=((G.get('alias') or '')+' · ' if G.get('alias') else '')+'Rund Alsaleh'
# ---------- 4) DB: new records (GitHub open item)
NEW=[
 ('أبرار محضار','Abrar Mihdhar','معسكر T5 لعلوم البيانات والذكاء الاصطناعي — أكاديمية سدايا','يناير 2022 (تاريخ مستودع مشروع المعسكر)','ds',72,'','',
  [['GitHub','https://github.com/Abrar-Mihdhar'],['مستودع المشروع','https://github.com/Abrar-Mihdhar/final-project-of-sdaiaT5_data-science-course']],
  'مستودع «final-project-of-sdaiaT5_data-science-course» (يناير 2022). الاسم من اسم المستخدم على GitHub (Abrar-Mihdhar)؛ ملف LinkedIn بالاسم بلا خبرات منشورة — لم يُنسب.'),
 ('سعود البلاع','Saud AlBallaa','معسكر T5 لعلوم البيانات والذكاء الاصطناعي — أكاديمية سدايا','يناير 2022 (تاريخ مستودع مشروع المعسكر)','ds',75,'','',
  [['GitHub','https://github.com/Saud-AlBallaa'],['مستودع المشروع','https://github.com/Saud-AlBallaa/Sdaia_Project']],
  'مستودع «Sdaia_Project — This is the final project submission» (يناير 2022)، والاسم الكامل على ملف GitHub (الرياض). ملف LinkedIn محتمل بالاسم نفسه (Senior Database Engineer — بنك البلاد) لا يُظهر قيد سدايا — لم يُنسب.'),
 ('هند العسكر','Hind Alaskar','معسكر T5 لعلوم البيانات والذكاء الاصطناعي — أكاديمية سدايا','يناير 2022 (تاريخ مستودع المعسكر)','ds',75,'','',
  [['GitHub','https://github.com/HindAlaskar'],['مستودع المعسكر','https://github.com/HindAlaskar/T5SDAIA']],
  'مستودع «T5SDAIA» (يناير 2022)، والاسم الكامل على ملف GitHub (الرياض).'),
 ('هناء العمري','Hana Alamri','معسكر T5 لعلوم البيانات والذكاء الاصطناعي — أكاديمية سدايا','يناير 2022 (تاريخ مستودع المعسكر)','ds',70,'','',
  [['GitHub','https://github.com/hana-alamri'],['مستودع المعسكر','https://github.com/hana-alamri/T5SDAIA']],
  'مستودع «T5SDAIA» (يناير 2022)؛ الاسم من اسم المستخدم على GitHub (hana-alamri).'),
 ('هبة الجاسر','Heba AlJassir','معسكر T5 لعلوم البيانات والذكاء الاصطناعي — أكاديمية سدايا','يناير 2022 (تاريخ مستودع مشروع المعسكر)','ds',70,'','',
  [['GitHub','https://github.com/HebaAlJassir'],['مستودع المشروع','https://github.com/HebaAlJassir/T5-Bootcamp-Project']],
  'مستودع «T5-Bootcamp-Project» (يناير 2022، توقيت دفعة T5 الثانية)؛ الاسم من اسم المستخدم على GitHub.'),
 ('مصعب البرقي','Mussab Albargi','معسكر T5 لعلوم البيانات والذكاء الاصطناعي — أكاديمية سدايا','يناير 2024 (تاريخ مستودع المعسكر)','ds',72,'','',
  [['GitHub','https://github.com/Mussab-albargi'],['مستودع المعسكر','https://github.com/Mussab-albargi/SDAIA_T5_AI_BOOTCAMP']],
  'مستودع «SDAIA_T5_AI_BOOTCAMP» (يناير 2024)؛ الاسم من اسم المستخدم على GitHub.'),
 ('الأقصى أكبر','Alaqsa Akbar','معسكر الذكاء الاصطناعي سدايا–كاوست (SDAIA-KAUST AI Bootcamp)','فبراير 2024 (تاريخ مستودع المعسكر)','ai',78,'هيوماين (HUMAIN)','مهندس ذكاء اصطناعي (يونيو 2026)',
  [['GitHub','https://github.com/alaqsa-akbar'],['مستودع المعسكر','https://github.com/alaqsa-akbar/KAUST_SDAIA_AI_Bootcamp_1'],['LinkedIn','https://www.linkedin.com/in/alaqsa-akbar/']],
  'مستودع «KAUST_SDAIA_AI_Bootcamp_1» (فبراير 2024)؛ ملف GitHub «AI Engineer @ HUMAIN» يطابق صفحة الخبرة على LinkedIn: Artificial Intelligence Engineer — HUMAIN · Full-time منذ يونيو 2026 (+28، خارج نافذة القصص).'),
]
new_ids=[]
sys.path.insert(0,H+'Graduates/tools'); import translit as TR
idx=TR.build_index(g)
for ar,en,prog,co,tr,score,emp,role,links,d in NEW:
    m=TR.match(ar,en,idx)
    assert not m,(en,m)
    rid=len(g)+1; r=rid+1
    cats=['program','project']+(['job'] if emp else [])
    dd=d+' '+TAG.strip()+TL
    G={"n":ar,"en":en,"lv":"gh1007","score":score,"prog":prog,"co":co,"win":"ضمن النطاق","edu":"","emp":emp,"role":role,"cats":cats,"d":dd,"links":links,"src":SRC,"tr":tr}
    g.append(G); new_ids.append(rid)
    li=' · '.join(l[1] for l in links if 'linkedin' in l[1]) or None
    gh=' · '.join(l[1] for l in links if 'github' in l[1]) or None
    res='برنامج · مشروع'+(' · توظيف' if emp else '')
    vals=[rid,ar,en,prog,TRN[tr],co,emp or None,role or None,None,None,'gh1007',score,'ضمن النطاق',res,li,gh,None,None,None,None,dd,None,SRC]
    for ci,v in enumerate(vals,1): ws.cell(r,ci).value=v
# ---------- 5) STORIES
S=[x for x in ST if x['id']!='amal-almubarak']; assert len(S)==167
SI={x['id']:x for x in S}
GAP_RE=re.compile(r' ?⚠ ملاحظة تحرير:.*$')
def retail(sid,txt):
    x=SI[sid]; st=x['story']; assert '⚠ ملاحظة تحرير' in st,sid
    keep=''
    if sid=='yasser-alshehri': keep=' وتجدر الإشارة إلى وجود أكثر من سجل باسم مشابه في قاعدة الخريجين؛ رُبطت القصة بالسجل الأقرب تطابقًا في الاسم والمسار.'
    x['story']=GAP_RE.sub('',st).rstrip()+' '+txt+keep
PROVT=f' + تأريخ شهري من قسمي الشهادات والخبرة على LinkedIn بحساب مسجَّل (قُرئ {DATE}).'
U={
 'amal-almuarik':dict(period='نوفمبر 2024 – مارس 2025',year=2025,clear=True,
   tail='التأريخ الشهري (صفحة الخبرة على LinkedIn): معسكر Microsoft لهندسة البيانات — SDAIA (نوفمبر 2024 – مارس 2025)، ثم برنامج تطوير الخريجين Elite (مايو–ديسمبر 2025)، ثم أخصائية ذكاء اصطناعي بدوام كامل في GOSI منذ ديسمبر 2025 — بعد تسعة أشهر من الإتمام، داخل نافذة الأربعة عشر شهرًا.'),
 'bushra-dajam':dict(period='يونيو 2025',year=2025,clear=True,
   tail='التأريخ الشهري (قسما الشهادات والخبرة على LinkedIn): شهادتا NVIDIA-Certified Associate في النماذج اللغوية الكبيرة ومتعدد الوسائط (يونيو 2025)، ثم مهندسة ذكاء اصطناعي بدوام كامل في Arabot منذ ديسمبر 2025 — بعد ستة أشهر، داخل نافذة الأربعة عشر شهرًا.'),
 'nasser-alkuhili':dict(period='فبراير 2025',year=2025,clear=True,
   tail='التأريخ الشهري (قسما الشهادات والخبرة على LinkedIn): شهادتا NVIDIA-Certified Associate (فبراير 2025)، ثم مهندس ذكاء اصطناعي وكيلي بدوام كامل في IBM (أكتوبر 2025، +8)، ثم مهندس ذكاء اصطناعي في وحدة ابتكار الذكاء الاصطناعي بمصرف الراجحي منذ فبراير 2026 (+12) — كلا التعيينين داخل نافذة الأربعة عشر شهرًا.'),
 'lama-alzahrani':dict(period='نوفمبر 2024 – أبريل 2025',year=2025,clear=True,
   tail='التأريخ الشهري (صفحة الخبرة على LinkedIn): معسكر هندسة البيانات — SDAIA (نوفمبر 2024 – أبريل 2025)، ثم محللة تسويق بدوام كامل في stc منذ سبتمبر 2025 — بعد خمسة أشهر من الإتمام.'),
 'aljohara-alhogail':dict(period='أكتوبر – ديسمبر 2025',year=2025,clear=True,
   tail='التأريخ الشهري (صفحة الخبرة على LinkedIn): قيد «Cloud Systems Management Bootcamp — SDAIA» (أكتوبر–ديسمبر 2025، مع تكريم التخرج في ICAN 2025)، ثم مطوّرة RPA بدوام كامل في مصرف الإنماء منذ يناير 2026 — بعد شهر واحد من الإتمام. تسمّي قائمة الأكاديمية البرنامج «وكلاء الذكاء الاصطناعي»، وتسمّيه صفحة الخبرة «إدارة الأنظمة السحابية».',
   ach_add='معسكر إدارة الأنظمة السحابية — أكاديمية سدايا (أكتوبر–ديسمبر 2025)'),
 'etab-alotaibi':dict(period='فبراير 2024 · فبراير 2025',year=2025,clear=True,
   tail='التأريخ الشهري (قسما الشهادات والخبرة على LinkedIn): شهادة معسكر النماذج اللغوية الكبيرة — SDAIA (5 فبراير 2024)، ثم شهادتا NVIDIA في النماذج اللغوية الكبيرة ومتعدد الوسائط (فبراير 2025)، ثم أخصائية ذكاء اصطناعي بدوام كامل في كيمارك منذ أبريل 2026 — بعد أربعة عشر شهرًا من الإتمام الثاني، داخل النافذة الشاملة لشهر النهاية.'),
 'rand-alsaleh':dict(period='2024',year=2024,clear=True,link=('LinkedIn','https://www.linkedin.com/in/rundalsaleh/'),
   tail='التأريخ الشهري (صفحة الخبرة على LinkedIn باسم «Rund Alsaleh»): محللة تحوّل رقمي بدوام كامل في صدارة منذ فبراير 2025. ولأن الإتمام في 2024 بحسب قائمة الأكاديمية، فإن أبعد فارق ممكن (يناير 2024 ← فبراير 2025) ثلاثة عشر شهرًا — داخل النافذة أيًّا كان شهر الإتمام.'),
 'yasser-alshehri':dict(period='يناير – يوليو 2025',year=2025,clear=True,
   tail='التأريخ الشهري (قسما الشهادات والخبرة على LinkedIn): دورات NVIDIA المعتمدة (يناير–يوليو 2025) ومعسكر المعالجة الآلية للغة العربية ANLP — SDAIA (مايو 2025)، ثم أخصائي ذكاء اصطناعي بدوام كامل في جهة حكومية (Confidential Government) منذ يناير 2026 — بعد ستة أشهر.',
   ach_add='معسكر المعالجة الآلية للغة العربية (ANLP) — سدايا (مايو 2025)'),
 'rawabi-allihyani':dict(period='نوفمبر 2024 – أبريل 2025',year=2025,clear=False,
   tail='⚠ ملاحظة تحرير (محدّثة 7 أكتوبر 2026): تُظهر صفحة الخبرة على LinkedIn معسكر هندسة البيانات (نوفمبر 2024 – أبريل 2025)، ثم «محلل بيانات» في أمانة العاصمة المقدسة منذ مايو 2026 (+13) بصيغة تدريب منتهٍ بالتوظيف (Apprenticeship) لا وظيفة بدوام كامل؛ يبقى وسم المراجعة والنشر قائم على اعتماد الأكاديمية المباشر.'),
}
for sid,u in U.items():
    x=SI[sid]; retail(sid,u['tail']); x['period']=u['period']; x['year']=u['year']
    if u['clear']: x.pop('reviewFlag',None); C['flag_cleared']+=1
    if u.get('ach_add'): x['achievements'].insert(1,u['ach_add'])
    if u.get('link'): x['links'].append({'label':u['link'][0],'url':u['link'][1]})
    x['provenance']=(x.get('provenance') or '')+PROVT
x=SI['badr-alshamrani']
x['story']=rep(x['story'],'أتمّ بدر برنامج وكلاء الذكاء الاصطناعي بأكاديمية سدايا في 2025، والتحق في السنة نفسها مهندسَ','أتمّ بدر برنامج وكلاء الذكاء الاصطناعي بأكاديمية سدايا (سبتمبر 2025)، والتحق في ديسمبر 2025 — بعد ثلاثة أشهر — مهندسَ',1)
x['story']=rep(x['story'],'ووقوع التوظيف داخل سنة إتمام البرنامج يستوفي معيار الأربعة عشر شهرًا المعتمد.','والتعيين بدوام كامل مؤرخ بصفحة الخبرة على LinkedIn، داخل معيار الأربعة عشر شهرًا.',1)
x['achievements']=[a if a!='توظيف في سنة إتمام البرنامج نفسها (2025)' else 'توظيف بدوام كامل بعد ثلاثة أشهر من الإتمام (ديسمبر 2025)' for a in x['achievements']]
x['provenance']+=f' تأكيد {DATE}: التعيين Full-time منذ ديسمبر 2025 (صفحة الخبرة).'; C['badr_dated']=1
x=SI['nujud-senan']
x['achievements'].append('مديرة مشروع — وزارة الاتصالات وتقنية المعلومات (Full-time منذ مايو 2025)')
x['provenance']+=f' مراجعة {DATE}: تبقى الجائزة النتيجة المحتسبة (الفئة: جوائز)، وأُضيف تعيين الوزارة (مايو 2025، +7) إنجازًا مؤرخًا ليتسق البطاقة مع نص القصة.'
def nk(t): return re.sub(r'[^a-z]','',(t or '').lower().replace('abdul','abd'))
def ar(t): return re.sub('[إأآا]','ا',t).replace('ة','ه').replace('ى','ي').replace(' ','')
IDX=collections.defaultdict(set)
for k,G_ in enumerate(g,1):
    for l in G_['links']: IDX[re.sub(r'https?://[a-z]+\.','',l[1]).rstrip('/').lower()].add(k)
    IDX['en:'+nk(G_['en'])].add(k); IDX['ar:'+ar(G_['n'])].add(k)
FIX={'nasser-alsaqer':[27],'khalid-alduwaysan':[1]}
pmiss=[]; PMAP={}
for x in S:
    if x.get('provenance'): continue
    hits=set()
    for l in x['links']:
        if re.search(r'github\.com/[^/]+/[^/]+',l['url']): continue
        hits|=IDX.get(re.sub(r'https?://[a-z]+\.','',l['url']).rstrip('/').lower(),set())
    if not hits: hits|=IDX.get('en:'+nk(x['nameEn']),set())|IDX.get('ar:'+ar(x['name']),set())
    if len(hits)>1:
        h2={k for k in hits if nk(g[k-1]['en'])==nk(x['nameEn']) or ar(g[k-1]['n'])==ar(x['name'])}
        h3={k for k in h2 if x['org'][:6] in (g[k-1]['emp'] or '')}
        hits=h3 or h2 or hits
    hits=sorted(FIX.get(x['id'],hits))
    if not hits: pmiss.append(x['id']); continue
    G_=g[hits[0]-1]; PMAP[x['id']]=hits
    ref=' / '.join('#%d'%k for k in hits)
    lk=' · '.join(l['url'] for l in x['links'][:3]) or 'قائمة الأكاديمية (لا رابط علني)'
    x['provenance']=f"سجل قاعدة الخريجين {ref} — مصدر السجل: {G_['src']}. أدلة البطاقة: {lk}. (أُضيف حقل المصدر في جولة سد الفجوات، {DATE}.)"
    C['prov_filled']+=1
print('provenance unmapped',pmiss); print('multi-map',{k:v for k,v in PMAP.items() if len(v)>1})
N=len(g); T=collections.Counter(r['tr'] for r in g)
wl=sum(1 for r in g if r.get('links')); uu=len({l[1].rstrip('/') for r in g for l in (r.get('links') or [])})
S1=len(S); E1=89
print('N',N,dict(T),'wl',wl,'uu',uu,'S1',S1,'counts',dict(C),{k:len(v) for k,v in LOG.items()},'new',new_ids)
print('remaining: en',sum(1 for r in g if not r['en']),'co',sum(1 for r in g if not r['co'].strip()),'win?',sum(1 for r in g if r['win']=='غير محدد'),'flags',sum(1 for x in S if x.get('reviewFlag')),'noprov',sum(1 for x in S if not x.get('provenance')),'nolinks',sum(1 for x in S if not x['links']),'nophoto',sum(1 for x in S if not x['photo']),'nologo',sum(1 for x in S if not x['orgLogo']))
json.dump(dict(N=N,T=T,wl=wl,uu=uu,new=new_ids,S1=S1,C=C,LOG=LOG),open('stats.json','w'),ensure_ascii=False)
if DRY: sys.exit(0)
# =============== WRITE ===============
TOT=format(N,',')
OLD=dict(N=3309,TOT='3,309',S=168,E=89,ai=1576,ds=950,dmg=330,other=193,genai=249,coop=11,wl=1373,uu='2,229')
VA_O='إصدار 6 أكتوبر 2026 (جولة المسح العميق والتنظيف النهائي'; VA_N='إصدار 7 أكتوبر 2026 (جولة سد الفجوات واستكمال البيانات'
VE_O='version 6 October 2026 (deep sweep and final cleaning round'; VE_N='version 7 October 2026 (gap-closing and data-completion round'
def common(t):
    t=t.replace(OLD['TOT'],TOT).replace(VA_O,VA_N).replace(VE_O,VE_N)
    for a,b in ((f"({OLD['S']} قصة)",f'({S1} قصة)'),(f"({OLD['S']} stories)",f'({S1} stories)'),(f"{OLD['S']} قصة نجاح موثّقة",f'{S1} قصة نجاح موثّقة'),(f"{OLD['S']} verified success stories",f'{S1} verified success stories'),(f"{OLD['S']} قصة",f'{S1} قصة'),(f"{OLD['S']} stories",f'{S1} stories')):
        t=t.replace(a,b)
    return t
head=common(head); tail=common(tail)
tail=rep(tail,"{ value: %d, label: 'سجلًا فرديًا',"%OLD['N'],"{ value: %d, label: 'سجلًا فرديًا',"%N,1)
for k,en in (('ai','AI Engineer'),('ds','Data Scientist'),('dmg','Data Management & Governance'),('other','Other / Specialized')):
    if T[k]!=OLD[k]: tail=rrep(tail,r"(en: '%s', value: )%d\b"%(re.escape(en),OLD[k]),"\\g<1>%d"%T[k])
tail=rrep(tail,r"\{ value: %d,(\s+)label: 'سجلًا برابط تحقق علني', note: '%s رابط مصدر فريد',"%(OLD['wl'],OLD['uu']),"{ value: %d,\\1label: 'سجلًا برابط تحقق علني', note: '%s رابط مصدر فريد',"%(wl,format(uu,',')))
tail=rep(tail,"noteEn: '%s unique source links'"%OLD['uu'],"noteEn: '%s unique source links'"%format(uu,','),1)
tail=rrep(tail,r"\{ value: %d,(\s+)label: 'قصة نجاح موثّقة'"%OLD['S'],"{ value: %d,\\1label: 'قصة نجاح موثّقة'"%S1)
D['grads']=g
open(P,'w',encoding='utf-8').write(head+json.dumps(D,ensure_ascii=False)+';'+tail)
print('grad leftovers',{k:(head+tail).count(k) for k in [OLD['TOT'],f"{OLD['S']} قصة",VA_O]})
P2=H+'index.html'; s=open(P2,encoding='utf-8').read(); s=common(s)
s=rep(s,'data-count="%d"'%OLD['N'],'data-count="%d"'%N,1); s=rep(s,'data-count="%d"'%OLD['S'],'data-count="%d"'%S1,1)
for k in ('ai','ds','dmg','other','genai','coop'):
    pc='%.1f'%(100*T[k]/N)
    pat=r'(data-count=")%d(">0</div>.*?--w:)[\d.]+(%%.*?data-ar=")[\d.]+(%% من السجل" data-en=")[\d.]+(%% of the registry")'%OLD[k]
    s,n=re.subn(pat,lambda m:m.group(1)+str(T[k])+m.group(2)+pc+m.group(3)+pc+m.group(4)+pc+m.group(5),s,count=1,flags=re.S); assert n==1,k
open(P2,'w',encoding='utf-8').write(s)
print('portal leftovers',{k:s.count(k) for k in [OLD['TOT'],f"{OLD['S']} قصة"]})
P3=H+'SuccessStories/website/js/data.js'; s=open(P3,encoding='utf-8').read()
s=rep(s,"{ value: %d, label: 'خريجًا موثّقًا',"%OLD['N'],"{ value: %d, label: 'خريجًا موثّقًا',"%N,1)
for k,en in (('ai','AI Engineer'),('ds','Data Scientist'),('dmg','Data Management & Governance'),('other','Other / Specialized')):
    if T[k]!=OLD[k]: s=rrep(s,r"(en: '%s',\s+value: )%d\b"%(re.escape(en),OLD[k]),"\\g<1>%d"%T[k])
s=rrep(s,r"(\{ value: )%d(,\s+label: 'قصة نجاح)"%OLD['S'],r"\g<1>%d\2"%S1)
k=s.index('const STORIES = [')+len('const STORIES = ['); e=s.index('\n];',k)
s=s[:k]+'\n'+',\n'.join('  '+json.dumps(x,ensure_ascii=False) for x in S)+s[e:]
s=rep(s,"'use strict';","/* %s: جولة سد الفجوات — %s خريجًا (%d سجلات جديدة)؛ القصص 168 ← %d (دمج قصة مكررة amal-almubarak)؛ %d وسم مراجعة أُزيل بتأريخ شهري؛ %d حقل مصدر أُكمل. */\n'use strict';"%(DATE,TOT,len(new_ids),S1,C['flag_cleared'],C['prov_filled']),1)
open(P3,'w',encoding='utf-8').write(s)
P4=H+'SuccessStories/website/index.html'; s=open(P4,encoding='utf-8').read(); s=rep(s,'"numberOfItems": %d'%OLD['S'],'"numberOfItems": %d'%S1,1); open(P4,'w',encoding='utf-8').write(s)
hi=sum(1 for r in g if r.get('score',0)>=75); mid=sum(1 for r in g if 45<=r.get('score',0)<75); lo=N-hi-mid
emp=sum(1 for r in g if r.get('emp'))
yc=collections.Counter()
for r in g:
    ys=set(re.findall(r'20(2[1-6])',r.get('co') or ''))
    if not ys: yc['none']+=1
    for y in ys: yc['20'+y]+=1
st=wb['إحصاءات']
M={'الإجمالي':TOT,'مهندس ذكاء اصطناعي':T['ai'],'عالم بيانات':T['ds'],'أكاديمية الذكاء الاصطناعي التوليدي':T['genai'],'إدارة وحوكمة البيانات':T['dmg'],'التدريب التعاوني':T['coop'],'برامج متخصصة أخرى':T['other'],
   'ثقة مرتفعة (≥75)':hi,'ثقة متوسطة (45–74)':mid,'ثقة منخفضة (<45)':lo,'سجلات بجهة عمل موثّقة':emp,'سجلات بروابط تحقق علنية':wl,'روابط مصدر فريدة':uu,
   'سجلات بدفعة 2021':yc['2021'],'سجلات بدفعة 2022':yc['2022'],'سجلات بدفعة 2023':yc['2023'],'سجلات بدفعة 2024':yc['2024'],'سجلات بدفعة 2025':yc['2025'],'سجلات بدفعة 2026':yc['2026'],'سجلات بلا دفعة محددة':yc['none'],
   'قصص النجاح':S1,'جهات العمل':E1,
   'تحديث':f'{DATE} — جولة سد الفجوات: {len(new_ids)} سجلات جديدة (3,309 → {TOT})؛ القصص 168 → {S1} (دمج قصة مكررة)؛ {C["flag_cleared"]} وسم مراجعة أُزيل؛ {C["co_set"]} دفعة و{C["en_set"]} اسمًا إنجليزيًا و{C["win_set"]} نطاقًا زمنيًا أُكملت'}
for row in st.iter_rows():
    if row[0].value in M: row[1].value=M[row[0].value]
sh=wb.create_sheet('جولة_سد_الفجوات_2026-10-07')
sh.append(['#','الاسم','النوع','التفصيل'])
for k in new_ids: sh.append([k,g[k-1]['n'],'سجل جديد (GitHub)',g[k-1]['prog']])
for k,(co,ev) in F.CO.items(): sh.append([k,g[k-1]['n'],'تأريخ الدفعة' if co else 'تأكيد التأريخ',ev])
for k in F.EN: sh.append([k,g[k-1]['n'],'اسم إنجليزي',F.EN[k]])
for k in LOG['win']: sh.append([k,g[k-1]['n'],'النطاق الزمني ← ضمن النطاق',g[k-1]['co']])
for k in LOG['staff']: sh.append([k,g[k-1]['n'],'⚠ مرشح للاستبعاد (موظف/لا دليل)',''])
for k,lab,url,txt in ENR: sh.append([k,g[k-1]['n'],'إثراء',txt[:200]])
wb.save(X_)
P5=H+'Graduates/Graduates_Database.md'; s=open(P5,encoding='utf-8').read()
s=rep(s,'**الإصدار:** 6 أكتوبر 2026 — ','**الإصدار:** %s — جولة سد الفجوات واستكمال البيانات (3,309 → %s · القصص 168 → %d). قبله: 6 أكتوبر 2026 — '%(DATE,TOT,S1),1)
L=['## جولة سد الفجوات واستكمال البيانات — %s'%DATE,'',f'**الخريجون 3,309 → {TOT} · القصص 168 → {S1} · جهات العمل {E1} (دون تغيير).**','',
 f'**1) تأريخ الدفعات ({len(F.CO)} سجلًا)** من قسمي الشهادات والخبرة على LinkedIn (حساب مسجَّل): {C["co_set"]} دفعة حُدّدت أو دُقّقت و{C["co_confirmed"]} تأكيد. ثم صُحّح «النطاق الزمني» من «غير محدد» إلى «ضمن النطاق» في {C["win_set"]} سجلًا مؤرخًا بين 2021–2022 و2024–2026.','',
 f'**2) أسماء إنجليزية ({C["en_set"]}):** من الاسم المعروض على LinkedIn أو GitHub — '+' · '.join(f'#{k} {v}' for k,v in F.EN.items())+'.','',
 f'**3) سجلات جديدة ({len(new_ids)})** — البند المفتوح من جولة 6 أكتوبر (مالكو مستودعات GitHub غير المدرجين): فُرز 62 مالكًا مرشحًا، واستُبعد من لا يظهر إلا باسم أول أو بالأحرف الأولى ومن لا يسمّي مستودعه سدايا/المعسكر ومن هو خارج المملكة.','',
 '| # | الاسم | الاسم الإنجليزي | البرنامج | الدفعة | الدليل |','|---|---|---|---|---|---|']
for k in new_ids:
    G_=g[k-1]; L.append(f"| {k} | {G_['n']} | {G_['en']} | {G_['prog']} | {G_['co']} | {G_['links'][1][1] if len(G_['links'])>1 else G_['links'][0][1]} |")
L+=['',f'**4) إثراءات:** '+' · '.join(sorted({f'#{k}' for k,_,_,_ in ENR}))+' (روابط GitHub/LinkedIn وتأريخ شهري).','',
 f'**5) مرشحون للاستبعاد ({len(LOG["staff"])}):** '+' · '.join(f'#{k}' for k in LOG['staff'])+' — موظفو سدايا ببرامج «مستنتجة» أو بلا أي دليل انتساب؛ وُسموا ولم يُحذفوا.','',
 f'**6) القصص:** دمج قصة مكررة (amal-almubarak ← amal-almuarik، الشخص نفسه)؛ إزالة وسم المراجعة `gap` عن {C["flag_cleared"]} قصص بتأريخ شهري داخل النافذة (amal-almuarik · bushra-dajam · nasser-alkuhili · lama-alzahrani · aljohara-alhogail · etab-alotaibi · rand-alsaleh · yasser-alshehri)؛ يبقى الوسم على rawabi-allihyani (تدريب منتهٍ بالتوظيف لا وظيفة بدوام كامل). تأريخ تعيين badr-alshamrani بالشهر (ديسمبر 2025)؛ إضافة رابط LinkedIn لـ rand-alsaleh؛ إكمال حقل المصدر في {C["prov_filled"]} قصة.','','---','']
k=s.index('## جولة المسح العميق والتنظيف النهائي — 6 أكتوبر 2026')
s=s[:k]+'\n'.join(L)+'\n'+s[k:]
open(P5,'w',encoding='utf-8').write(s)
print('done')
