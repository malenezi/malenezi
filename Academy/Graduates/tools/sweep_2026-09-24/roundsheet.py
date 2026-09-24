import openpyxl,json,os
from openpyxl.styles import Font,Alignment
P=os.path.expanduser('~/mnt/Academy/Graduates/Graduates_Database.xlsx')
wb=openpyxl.load_workbook(P); g=json.load(open('grads_after.json')); st=json.load(open('stats24.json')); S=json.load(open('sets.json'))
name='جولة_التسليم_2026-09-24'
if name in wb.sheetnames: del wb[name]
ws=wb.create_sheet(name); ws.sheet_view.rightToLeft=True
hdr=['#','الاسم العربي','الاسم الإنجليزي','البرنامج','الدفعة','المسار','الثقة','نوع العملية','حكم سياسة القصص','الدليل']
ws.append(hdr)
for c in ws[1]: c.font=Font(bold=True)
TRL={'ai':'مهندس ذكاء اصطناعي','ds':'عالم بيانات','other':'برامج متخصصة أخرى','dmg':'إدارة وحوكمة البيانات','genai':'أكاديمية الذكاء الاصطناعي التوليدي','coop':'التدريب التعاوني'}
ADJ={2735:'PASS_DATED — قصة marwah-ali',2832:'CERT_ONLY (تصحيح دفعة)',2750:'CERT_ONLY (تصحيح دفعة وجهة عمل)',2734:'لا نتيجة مؤرخة',2714:'FAIL_RELATEDNESS',362:'FAIL_TIMING',303:'HELD_ONE_DATUM',2700:'CERT_ONLY',2736:'لا نتيجة مؤهلة',2561:'CERT_ONLY'}
rows=[]
for idx in range(st['base']+1,len(g)+1):
    x=g[idx-1]; verdict='PASS_DATED — قصة meshal-aldalbahi' if x['en']=='Meshal Aldalbahi' else 'FAIL_COURSE_OUTPUT (§5-ب) — مخرَج مقرر، النافذة لم تنضج'
    rows.append([idx,x['n'],x['en'],x['prog'],x['co'],TRL[x['tr']],x['score'],'سجل جديد',verdict,x['links'][1][1] if len(x['links'])>1 else x['links'][0][1]])
for o,i in sorted(S['ENR'].items(),key=lambda t:t[1]):
    x=g[i-1]; rows.append([i,x['n'],x.get('en'),x['prog'],x['co'],TRL[x['tr']],x['score'],'إثراء بدل تكرار (GitHub: %s)'%o,'لا نتيجة لاحقة مؤرخة','https://github.com/'+o])
for i,v in ADJ.items():
    x=g[i-1]; rows.append([i,x['n'],x.get('en'),x['prog'],x['co'],TRL[x['tr']],x['score'],'حسم عبر صفحة الخبرة على LinkedIn',v,[u for l,u in x['links'] if 'linkedin' in u][0]])
for r in rows: ws.append(r)
for col,w in zip('ABCDEFGHIJ',[7,24,26,50,30,18,7,30,40,60]): ws.column_dimensions[col].width=w
ws.auto_filter.ref='A1:J%d'%(len(rows)+1)
wb.save(P); print(len(rows))
