import openpyxl,json,os,re,collections
from openpyxl.styles import Font
P=os.path.expanduser('~/mnt/Academy/Graduates/Graduates_Database.xlsx')
wb=openpyxl.load_workbook(P); ws=wb.worksheets[0]
g=json.load(open('grads_after.json')); st=json.load(open('stats24.json')); S=json.load(open('sets.json'))
TR={'ai':'مهندس ذكاء اصطناعي — AI Engineer','ds':'عالم بيانات — Data Scientist','dmg':'إدارة وحوكمة البيانات — Data Management & Governance','genai':'أكاديمية الذكاء الاصطناعي التوليدي — Gen AI Academy','other':'برامج متخصصة أخرى — Other / Specialized','coop':'التدريب التعاوني — Cooperative Training'}
CM={'cert':'شهادة','project':'مشروع','program':'برنامج','job':'توظيف','training':'تدريب','research':'بحث','lead':'قيادة','promo':'ترقية','award':'جائزة','startup':'ناشئة','upskill':'تطوير مهارات','story':'قصة','leadership':'قيادة'}
def outc(c):
    out=[]
    for k in c:
        v=CM.get(k,k)
        if v not in out: out.append(v)
    return ' · '.join(out) or None
def links(x):
    li=[u for l,u in x['links'] if 'linkedin.com/in/' in u]; gh=[u for l,u in x['links'] if re.match(r'https://github\.com/[^/]+/?$',u)]
    other=[u for l,u in x['links'] if u not in li[:1]+gh[:1]]
    return (li[0] if li else None),(gh[0] if gh else None),(' · '.join(other) or None)
changes=collections.Counter()
def setc(r,c,v,tag):
    if v=='' : v=None
    old=ws.cell(r,c).value
    if (str(old) if old is not None else None)!=(str(v) if v is not None else None) or type(old)!=type(v):
        ws.cell(r,c).value=v; changes[tag]+=1
for idx,x in enumerate(g,1):
    r=idx+1; newrow = idx>2973
    setc(r,1,idx,'id'); setc(r,2,x['n'],'n'); setc(r,3,x.get('en') or None,'en'); setc(r,4,x['prog'],'prog'); setc(r,5,TR[x['tr']],'track')
    setc(r,6,x.get('co') or None,'co'); setc(r,7,x.get('emp') or None,'emp'); setc(r,8,x.get('role') or None,'role'); setc(r,10,x.get('edu') or None,'edu')
    setc(r,11,x.get('lv'),'lv'); setc(r,12,int(x['score']),'score'); setc(r,13,x.get('win'),'win'); setc(r,14,outc(x['cats']),'outcomes')
    setc(r,21,x.get('d') or None,'d'); setc(r,23,x.get('src') or None,'src')
    li,gh,oth=links(x)
    if newrow:
        setc(r,15,li,'li'); setc(r,16,gh,'gh'); setc(r,20,oth,'other')
    else:
        if li and not ws.cell(r,15).value: setc(r,15,li,'li')
        if gh and not ws.cell(r,16).value: setc(r,16,gh,'gh')
        if idx in S['ENR'].values():
            cur=ws.cell(r,20).value or ''
            add=[u for l,u in x['links'] if 'github.com' in u and u not in cur and u!=ws.cell(r,16).value]
            if add: setc(r,20,(cur+' · ' if cur else '')+' · '.join(add),'other')
last=len(g)+1
ws.auto_filter.ref='A1:W%d'%last
print('changes',dict(changes))
# stats sheet
s2=wb['إحصاءات']; T=st['tracks']
vals={'الإجمالي':format(st['total'],','),'مهندس ذكاء اصطناعي':T['ai'],'عالم بيانات':T['ds'],'أكاديمية الذكاء الاصطناعي التوليدي':T['genai'],'إدارة وحوكمة البيانات':T['dmg'],'التدريب التعاوني':T['coop'],'برامج متخصصة أخرى':T['other'],
'ثقة مرتفعة (≥75)':st['hi'],'ثقة متوسطة (45–74)':st['md'],'ثقة منخفضة (<45)':st['lo'],'سجلات بجهة عمل موثّقة':st['we'],'سجلات بروابط تحقق علنية':st['wl'],'روابط مصدر فريدة':st['uu'],
'قصص النجاح':111,'جهات العمل':68}
for y,v in st['yrs'].items(): vals['سجلات بدفعة %s'%y]=v
for row in s2.iter_rows(min_row=1):
    k=row[0].value
    if k=='سجلات بدفعة NONE': row[0].value='سجلات بلا دفعة محددة'; k=row[0].value
    if k=='سجلات بلا دفعة محددة': row[1].value=st['none']
    elif k in vals: row[1].value=vals[k]
    elif k=='تحديث': row[1].value='24 سبتمبر 2026 — جولة مستودعات تسليم دفعات 2025–2026 (113 سجلًا جديدًا · 21 إثراءً · قصتا نجاح · تسوية أعمدة القاعدة مع DATA.grads)'
    elif k=='ترتيب الصفوف': row[1].value='مطابق تمامًا لترتيب DATA.grads في بوابة الخريجين — العمود «#» عدد صحيح هو المعرّف الحاكم (الصف = المعرّف + 1)'
json.dump(dict(changes),open('xlsx_changes.json','w'))
wb.save(P); print('saved', ws.max_row, ws.auto_filter.ref)
