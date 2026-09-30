import openpyxl,json,os,re,collections,shutil
from openpyxl.styles import Font,Alignment
P=os.path.expanduser('~/mnt/Academy/Graduates/Graduates_Database.xlsx')
b=P.replace('.xlsx','.backup-2026-09-30.xlsx')
if not os.path.exists(b): shutil.copy2(P,b)
wb=openpyxl.load_workbook(P); ws=wb.worksheets[0]
g=json.load(open('grads_after.json')); st=json.load(open('stats30.json')); ENR=set(st['enr'])|{1830,1999,2120,2992,1467,695,1476,1591,1822,1824}
assert ws.max_row==3086, ws.max_row
TR={'ai':'مهندس ذكاء اصطناعي — AI Engineer','ds':'عالم بيانات — Data Scientist','dmg':'إدارة وحوكمة البيانات — Data Management & Governance','genai':'أكاديمية الذكاء الاصطناعي التوليدي — Gen AI Academy','other':'برامج متخصصة أخرى — Other / Specialized','coop':'التدريب التعاوني — Cooperative Training'}
CM={'cert':'شهادة','project':'مشروع','program':'برنامج','job':'توظيف','training':'تدريب','research':'بحث','lead':'قيادة','promo':'ترقية','award':'جائزة','startup':'ناشئة','upskill':'تطوير مهارات','story':'قصة','leadership':'قيادة'}
def outc(c):
    out=[]
    for k in c:
        v=CM.get(k,k)
        if v not in out: out.append(v)
    return ' · '.join(out) or None
def links(x):
    li=[u for l,u in x['links'] if 'linkedin.com/in/' in u]; gh=[u for l,u in x['links'] if re.match(r'https://github\.com/[^/]+/?$',u)]
    other=[u for l,u in x['links'] if u not in li[:1]+gh[:1]]
    return (li[0] if li else None),(gh[0] if gh else None),(' · '.join(other) or None)
# sanity: existing rows match ids and names
for idx in (1,500,2000,3085):
    assert ws.cell(idx+1,1).value==idx and ws.cell(idx+1,2).value==g[idx-1]['n'] or idx in ENR,(idx,ws.cell(idx+1,1).value,ws.cell(idx+1,2).value)
ch=collections.Counter()
def setc(r,c,v,tag):
    if v=='': v=None
    if ws.cell(r,c).value!=v: ws.cell(r,c).value=v; ch[tag]+=1
for idx,x in enumerate(g,1):
    new=idx>3085
    if not new and idx not in ENR: continue
    r=idx+1
    if not new: assert ws.cell(r,1).value==idx and ws.cell(r,2).value==x['n'],(idx,ws.cell(r,2).value)
    setc(r,1,idx,'id'); setc(r,2,x['n'],'n'); setc(r,3,x.get('en') or None,'en'); setc(r,4,x['prog'],'prog'); setc(r,5,TR[x['tr']],'track')
    setc(r,6,x.get('co') or None,'co'); setc(r,7,x.get('emp') or None,'emp'); setc(r,8,x.get('role') or None,'role'); setc(r,10,x.get('edu') or None,'edu')
    setc(r,11,x.get('lv'),'lv'); setc(r,12,int(x['score']),'score'); setc(r,13,x.get('win'),'win'); setc(r,14,outc(x['cats']),'outcomes')
    setc(r,21,x.get('d') or None,'d'); setc(r,23,x.get('src') or None,'src')
    li,gh,oth=links(x)
    if new:
        setc(r,15,li,'li'); setc(r,16,gh,'gh'); setc(r,20,oth,'other')
    else:
        if li and not ws.cell(r,15).value: setc(r,15,li,'li')
        if gh and not ws.cell(r,16).value: setc(r,16,gh,'gh')
        cur=ws.cell(r,20).value or ''
        add=[u for l,u in x['links'] if u not in cur and u!=ws.cell(r,15).value and u!=ws.cell(r,16).value]
        if add: setc(r,20,(cur+' · ' if cur else '')+' · '.join(add),'other')
ws.auto_filter.ref='A1:W%d'%(len(g)+1)
# parity check
for idx,x in enumerate(g,1):
    assert ws.cell(idx+1,1).value==idx and ws.cell(idx+1,2).value==x['n'],idx
print('changes',dict(ch))
s2=wb['إحصاءات']; T=st['tracks']
vals={'الإجمالي':format(st['total'],','),'مهندس ذكاء اصطناعي':T['ai'],'عالم بيانات':T['ds'],'أكاديمية الذكاء الاصطناعي التوليدي':T['genai'],'إدارة وحوكمة البيانات':T['dmg'],'التدريب التعاوني':T['coop'],'برامج متخصصة أخرى':T['other'],
'ثقة مرتفعة (≥75)':st['hi'],'ثقة متوسطة (45–74)':st['md'],'ثقة منخفضة (<45)':st['lo'],'سجلات بجهة عمل موثّقة':st['we'],'سجلات بروابط تحقق علنية':st['wl'],'روابط مصدر فريدة':st['uu'],'قصص النجاح':112,'جهات العمل':69}
for y,v in st['yrs'].items(): vals['سجلات بدفعة %s'%y]=v
seen=set()
for row in s2.iter_rows(min_row=1):
    k=row[0].value
    if k=='سجلات بلا دفعة محددة': row[1].value=st['none']; seen.add(k)
    elif k in vals: row[1].value=vals[k]; seen.add(k)
    elif k=='تحديث': row[1].value='30 سبتمبر 2026 — جولة الإثراء السريع لدفعات 2024–2026 (95 سجلًا جديدًا: 54 من GitHub · 18 من LinkedIn · 23 من قوائم فرق مشاريع التخرج؛ 19 إثراءً؛ لا قصص جديدة)'; seen.add(k)
print('stats keys updated',len(seen),'missing',[k for k in vals if k not in seen])
name='جولة_إثراء_2026-09-30'
if name in wb.sheetnames: del wb[name]
w3=wb.create_sheet(name)
w3.sheet_view.rightToLeft=True
hdr=['#','النوع','الاسم العربي','الاسم الإنجليزي','البرنامج','الدفعة','المسار','الثقة','الطبقة','الروابط','الدليل']
w3.append(hdr)
for c in w3[1]: c.font=Font(bold=True)
for idx in range(3086,len(g)+1):
    x=g[idx-1]; w3.append([idx,'جديد',x['n'],x.get('en') or None,x['prog'],x['co'],TR[x['tr']],x['score'],x['lv'],' · '.join(u for l,u in x['links']) or None,x['d']])
for e in json.load(open('enr30.json')):
    x=g[e[0]-1]; w3.append([e[0],'إثراء',x['n'],x.get('en') or None,x['prog'],x['co'],TR[x['tr']],x['score'],x['lv'],' · '.join(u for l,u in x['links']) or None,e[2]])
for col,wd in zip('ABCDEFGHIJK',(7,8,22,24,50,34,26,7,11,60,90)): w3.column_dimensions[col].width=wd
wb.save(P); print('saved',ws.max_row,wb.sheetnames[-1])
