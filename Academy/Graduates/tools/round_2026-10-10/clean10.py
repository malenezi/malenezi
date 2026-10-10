# -*- coding: utf-8 -*-
import json,re,os,sys,collections
sys.path[:0]=[os.path.expanduser('~/mnt/Academy/Graduates/tools/round_2026-10-07')]
import openpyxl
H=os.path.expanduser('~/mnt/Academy/'); TAG='[تدقيق 10 أكتوبر 2026] '
P=H+'Graduates/index.html'; s=open(P,encoding='utf-8').read()
i=s.index('const DATA = ')+len('const DATA = '); j=s.index('\nconst SITE_STATS')
D=json.loads(s[i:j].rstrip()[:-1]); g=D['grads']; assert len(g)==3324
wb=openpyxl.load_workbook(H+'Graduates/Graduates_Database.xlsx'); ws=wb['قاعدة البيانات']
def note(k,t):
    g[k-1]['d']+=' '+TAG+t; r=k+1; assert ws.cell(r,1).value==k; ws.cell(r,21).value=(ws.cell(r,21).value or '')+' '+TAG+t
STRONG=[(2126,2400,'الاسم الكامل نفسه «Abdulaziz Alothman» في البرنامج نفسه والدفعة نفسها (AAI B5، ديسمبر 2025 – فبراير 2026): #2126 من سجل الأكاديمية + LinkedIn، و#2400 من مستودعات الأسابيع على GitHub (2xazo). صفتان مستقلتان متطابقتان (الاسم + الدفعة) — مكرر مرجّح؛ يُعرض للدمج (الإبقاء على #2126).'),
 (294,469,'الاسم الثلاثي نفسه «Asma Mubarak Albuainain» في قائمتين رسميتين (روستر NVIDIA و ML 9 نوفمبر 2023) — الشخص نفسه مرجّحًا في برنامجين؛ يُعرض للدمج (الإبقاء على #294).'),
 (301,676,'الاسم الثلاثي نفسه «Saad Awad Alghamdi» في قائمتين رسميتين (روستر NVIDIA و LLM BC 2023) — يُعرض للدمج (الإبقاء على #301).'),
 (348,504,'الاسم الثلاثي نفسه «Faten Naif Almutairi» في قائمتين رسميتين (روستر NVIDIA و LLM BC 2023) — يُعرض للدمج (الإبقاء على #348).'),
 (357,520,'الاسم الثلاثي نفسه «Ghadi Hassan Babour» في قائمتين رسميتين (روستر NVIDIA و LLM BC 2023) — يُعرض للدمج (الإبقاء على #357).')]
for a,b,t in STRONG: note(a,'⚠ مرشح دمج مع #%d: '%b+t); note(b,'⚠ مرشح دمج مع #%d: '%a+t)
WEAK=[(183,2869),(237,2749),(407,3080),(1602,3031),(1822,3160),(1900,3071),(1999,3113),(2105,2598),(2120,3111),(2192,2903),(2628,2884),(2873,3066),(2992,3137)]
for a,b in WEAK:
    assert g[a-1]['n'].strip()==g[b-1]['n'].strip(),(a,b)
    for x,y in ((a,b),(b,a)): note(x,f'إشارة متبادلة: الاسم نفسه «{g[x-1]["n"]}» في #{y} ({g[y-1]["prog"][:40]}) — اسم ثنائي شائع بلا رابط مشترك؛ لا دمج.')
s=s[:i]+json.dumps(D,ensure_ascii=False)+';'+s[j:]; open(P,'w',encoding='utf-8').write(s)
sh=wb['جولة_10_أكتوبر_2026']
for a,b,t in STRONG: sh.append([a,g[a-1]['n'],'⚠ مرشح دمج مع #%d'%b,t[:200]])
for a,b in WEAK: sh.append([a,g[a-1]['n'],'إشارة متبادلة (اسم مكرر) #%d'%b,'لا دمج'])
wb.save(H+'Graduates/Graduates_Database.xlsx'); print('ok',len(STRONG),len(WEAK))
